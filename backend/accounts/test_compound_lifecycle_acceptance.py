"""A06: real lifecycle/appointment/Marks combinations in disposable databases.

These tests catch partial closure, revived assignments, leaked terminal history,
lost audit atomicity and duplicate concurrent transitions. They do not certify UI.
"""
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
from threading import Barrier
from unittest.mock import patch

from django.apps import apps
from django.db import DatabaseError, close_old_connections, connection
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient, APITestCase, APITransactionTestCase
from rest_framework_simplejwt.token_blacklist.models import BlacklistedToken
from rest_framework_simplejwt.tokens import RefreshToken

from accounts.authorization import coordinator_programmes
from accounts.delegations import grant_delegation, revoke_delegation, today
from accounts.models import Coordinator, ParticipantLifecycleAudit, Student, User
from accounts.participant_lifecycle import (
    ParticipantLifecycleConflict, transition_lecturer, transition_student,
)
from appointments import co_supervision
from appointments import test_appointment_lifecycle as primary
from appointments.models import (
    CoSupervisorNomination, ResearchAmendment, SupervisorAppointment,
)
from appointments.test_co_supervision_lifecycle import submitted_marks
from marks.completion_windows import grant_completion_windows, task_access
from marks.models import EvaluationTask, MarkEntry, MarkScore, MarkSubmissionReceipt
from marks.services import ensure_period_tasks
from marks.test_write_helpers import submission_payload, with_version


class CompoundLifecycleFixture:
    _panel_appointment = primary.AppointmentLifecycleTests._panel_appointment

    def _lecturer(self, email, staff_no, *, supervisor=False, panel=False):
        return primary.AppointmentLifecycleTests._lecturer(
            self, email, staff_no, supervisor=supervisor or panel, panel=panel,
        )

    def setUp(self):
        self.assertTrue(str(connection.settings_dict['NAME']).startswith('test_'))
        primary.AppointmentLifecycleTests.setUp(self)
        self.assertFalse(self.office.is_staff or self.office.is_superuser)
        self.panel_appointment = self._panel_appointment()
        nomination = co_supervision.nominate(
            actor=self.supervisor, student_id=self.student.pk,
            candidate_id=self.new_supervisor.pk, justification='PRIVATE supporting expertise',
        )
        co_supervision.decide(nomination_id=nomination.pk, actor=self.new_supervisor, action='accept')
        co_supervision.decide(nomination_id=nomination.pk, actor=self.coordinator, action='approve')
        self.supporting = nomination.appointment
        self.entry, self.period = submitted_marks(self)

    def post(self, actor, path, data, expected=200):
        self.client.force_authenticate(actor)
        response = self.client.post(path, data, format='json')
        self.assertEqual(response.status_code, expected, getattr(response, 'data', None))
        return response

    def transition(self, target, reason='PRIVATE lifecycle decision', expected=200):
        return self.post(self.office,
            f'/api/accounts/participants/students/{self.student.matric_no}/transition/',
            {'targetStatus': target, 'reason': reason}, expected)

    def lecturer_transition(self, actor, target, expected=200):
        return self.post(self.office,
            f'/api/accounts/participants/lecturers/{actor.lecturer.staff_no}/transition/',
            {'targetStatus': target, 'reason': 'PRIVATE retirement decision'}, expected)

    def draft(self, evaluator=None, role='PANEL'):
        task = EvaluationTask.objects.create(
            profile=self.profile, period=self.period, evaluator=evaluator or self.panel,
            evaluator_role=role, assigned_by=self.office,
        )
        entry = MarkEntry.objects.create(task=task, comments='Saved draft survives', total_mark=46)
        MarkScore.objects.create(entry=entry, component=self.period.rubric.components.get(), marks_awarded=46)
        return task

    def window(self, task):
        grant_completion_windows(task_ids=[task.pk], actor=self.office,
            deadline=timezone.now() + timedelta(days=10), reason='Documented additional time')
        return task.completion_windows.get()

    def pending_support(self):
        return co_supervision.nominate(actor=self.supervisor, student_id=self.student.pk,
            candidate_id=self.new_panel.pk, justification='PRIVATE pending supporting expertise')

    def amendment(self):
        response = self.post(self.student_user, '/api/appointments/research-amendments/',
            {'kind': 'RESEARCH', 'title': 'Proposed research amendment', 'reason': 'Academic scope'}, 201)
        return ResearchAmendment.objects.get(pk=response.data['id'])

    def snapshot(self):
        labels = {'accounts', 'academics', 'appointments', 'marks', 'dashboard',
                  'announcements', 'token_blacklist'}
        return {model._meta.label: list(model.objects.order_by('pk').values())
                for model in apps.get_models() if model._meta.app_label in labels}

    def marks_snapshot(self):
        return {model: list(model.objects.order_by('pk').values()) for model in (MarkEntry, MarkScore)}

    def get(self, actor, path, expected=200):
        self.client.force_authenticate(actor)
        response = self.client.get(path)
        self.assertEqual(response.status_code, expected, getattr(response, 'data', None))
        return response

    def assert_closed(self, outcome):
        for appointment in (self.supervisor_appointment, self.panel_appointment, self.supporting):
            appointment.refresh_from_db()
            self.assertEqual((appointment.status, appointment.end_outcome), ('ENDED', outcome))
            self.assertEqual(appointment.lifecycle_events.filter(action='ENDED').count(), 1)


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class CompoundLifecycleAcceptanceTests(CompoundLifecycleFixture, APITestCase):
    def test_withdrawal_closes_team_cancels_support_and_amendment_preserving_marks(self):
        pending = self.pending_support()
        amendment = self.amendment()
        task = self.draft()
        window = self.window(task)
        before = self.marks_snapshot()
        self.transition('WITHDRAWN')
        self.assert_closed('WITHDRAWN')
        pending.refresh_from_db()
        amendment.refresh_from_db()
        task.refresh_from_db()
        window.refresh_from_db()
        self.assertEqual((pending.status, amendment.status), ('CANCELLED', 'CANCELLED'))
        self.assertEqual(task.lifecycle_status, 'RETIRED')
        self.assertEqual(window.end_kind, 'INVALIDATED')
        self.assertEqual(self.marks_snapshot(), before)
        self.assertEqual(task.lifecycle_audits.get().entry_snapshot['comments'], 'Saved draft survives')
        self.assertEqual(self.student.lifecycle_audits.count(), 1)
        self.assertFalse(EvaluationTask.objects.filter(evaluator=self.new_supervisor).exists())

    def test_graduation_requires_supporting_nomination_resolution_then_completes_team(self):
        pending = self.pending_support()
        before = self.snapshot()
        blocked = self.transition('GRADUATED', expected=409)
        self.assertEqual(blocked.data['blockers']['pendingCoSupervisorNominations'], 1)
        self.assertEqual(self.snapshot(), before)
        self.post(self.office,
            f'/api/accounts/participants/students/{self.student.matric_no}/pending-work/cancel/',
            {'recordType': 'CO_SUPERVISOR_NOMINATION', 'recordId': pending.pk,
             'reason': 'No further supporting appointment needed'})
        marks = self.marks_snapshot()
        self.transition('GRADUATED')
        self.assert_closed('COMPLETED')
        self.assertEqual(self.marks_snapshot(), marks)
        self.transition('ACTIVE', expected=409)

    def test_graduation_with_unfinished_marks_is_atomic_even_with_supporting_team(self):
        self.draft()
        before = self.snapshot()
        blocked = self.transition('GRADUATED', expected=409)
        self.assertEqual(blocked.data['blockers']['unfinishedMarksTasks'], 1)
        self.assertEqual(self.snapshot(), before)

    def test_reactivation_after_panel_ended_retires_old_draft_without_restoring_window(self):
        task = self.draft()
        window = self.window(task)
        marks = self.marks_snapshot()
        self.transition('DEFERRED')
        self.post(self.office,
            f'/api/appointments/panel/appointments/{self.panel_appointment.pk}/end/',
            {'outcome': 'OTHER', 'reason': 'Panel appointment ended during leave'})
        self.transition('ACTIVE', 'Returned after panel departure')
        task.refresh_from_db()
        window.refresh_from_db()
        self.assertEqual(task.lifecycle_status, 'RETIRED')
        self.assertEqual(window.end_kind, 'INVALIDATED')
        self.assertFalse(task_access(task)['canEdit'])
        self.assertEqual(list(task.lifecycle_audits.order_by('created_at', 'pk').values_list('action', flat=True)),
                         ['PAUSED', 'RETIRED'])
        self.assertEqual(self.marks_snapshot(), marks)
        ensure_period_tasks(self.period, actor=self.office)
        self.assertFalse(EvaluationTask.objects.filter(profile=self.profile, evaluator=self.panel,
            lifecycle_status='ACTIVE').exists())
        self.client.force_authenticate(self.panel)
        response = self.client.put(f'/api/marks/tasks/{task.pk}/draft/',
            with_version(task, {'comments': 'Cannot revive ended assignment'}), format='json')
        self.assertEqual(response.status_code, 404, response.data)

    def test_reactivation_resumes_current_team_but_never_restores_old_window(self):
        task = self.draft()
        window = self.window(task)
        marks = self.marks_snapshot()
        self.transition('DEFERRED')
        self.transition('ACTIVE')
        task.refresh_from_db()
        window.refresh_from_db()
        self.assertEqual(task.lifecycle_status, 'ACTIVE')
        self.assertEqual(window.end_kind, 'INVALIDATED')
        self.assertTrue(task_access(task)['canEdit'])
        self.assertEqual(self.marks_snapshot(), marks)
        self.assertEqual(self.supporting.status, 'ACTIVE')

    def test_withdrawal_audit_outage_rolls_back_every_team_task_window_and_amendment(self):
        self.pending_support()
        self.amendment()
        self.window(self.draft())
        before = self.snapshot()
        with patch('accounts.participant_lifecycle.ParticipantLifecycleAudit.objects.create',
                   side_effect=DatabaseError('Controlled audit write outage')):
            with self.assertRaises(DatabaseError):
                transition_student(matric_no=self.student.matric_no, actor=self.office,
                    target_status='WITHDRAWN', reason='Withdrawal must be atomic')
        self.assertEqual(self.snapshot(), before)
        self.transition('WITHDRAWN')
        self.assert_closed('WITHDRAWN')

    def test_marks_audit_outage_rolls_back_deferral_and_completion_window(self):
        self.window(self.draft())
        before = self.snapshot()
        with patch('marks.services.EvaluationTaskLifecycleAudit.objects.create',
                   side_effect=DatabaseError('Controlled Marks audit outage')):
            with self.assertRaises(DatabaseError):
                transition_student(matric_no=self.student.matric_no, actor=self.office,
                    target_status='DEFERRED', reason='Deferral must be atomic')
        self.assertEqual(self.snapshot(), before)

    def test_coordinator_retirement_requires_managed_programme_and_delegation_resolution(self):
        delegation = grant_delegation(actor=self.office, coordinator_id=self.other_coordinator.pk,
            programme=self.student.programme, starts_on=today(), ends_on=today() + timedelta(days=2),
            justification='Temporary programme coverage')
        self.lecturer_transition(self.other_coordinator, 'RETIRING')
        blocked = self.lecturer_transition(self.other_coordinator, 'RETIRED', 409)
        self.assertEqual(blocked.data['blockers']['managedProgrammes'], 1)
        self.assertEqual(blocked.data['blockers']['coordinatorDelegations'], 1)
        self.assertEqual(coordinator_programmes(self.other_coordinator), [])
        revoke_delegation(delegation_id=delegation.pk, actor=self.office, reason='Coverage reassigned')
        blocked = self.lecturer_transition(self.other_coordinator, 'RETIRED', 409)
        self.assertEqual(blocked.data['blockers']['managedProgrammes'], 1)
        self.assertEqual(blocked.data['blockers']['coordinatorDelegations'], 0)
        # Controlled setup of an already completed responsibility handover.
        Coordinator.objects.filter(pk=self.other_coordinator.pk).update(programme_managed='')
        before = self.marks_snapshot()
        self.lecturer_transition(self.other_coordinator, 'RETIRED')
        self.other_coordinator.refresh_from_db()
        self.assertFalse(self.other_coordinator.is_active)
        self.assertEqual(self.marks_snapshot(), before)
        self.assertEqual(self.other_coordinator.lecturer.lifecycle_audits.count(), 2)

    def test_retirement_audit_outage_preserves_login_refresh_token_and_draft(self):
        task = self.draft(self.new_panel, 'BACKUP')
        self.window(task)
        self.lecturer_transition(self.new_panel, 'RETIRING')
        token = RefreshToken.for_user(self.new_panel)
        before = self.snapshot()
        with patch('accounts.participant_lifecycle.ParticipantLifecycleAudit.objects.create',
                   side_effect=DatabaseError('Controlled retirement audit outage')):
            with self.assertRaises(DatabaseError):
                transition_lecturer(staff_no=self.new_panel.lecturer.staff_no, actor=self.office,
                    target_status='RETIRED', reason='Retirement must be atomic')
        self.assertEqual(self.snapshot(), before)
        self.lecturer_transition(self.new_panel, 'RETIRED')
        task.refresh_from_db()
        self.new_panel.refresh_from_db()
        self.assertEqual(task.lifecycle_status, 'RETIRED')
        self.assertFalse(self.new_panel.is_active)
        self.assertTrue(BlacklistedToken.objects.filter(token__jti=token['jti']).exists())

    def test_retiring_supporting_candidate_pending_work_must_be_cancelled_before_retirement(self):
        pending = self.pending_support()
        self.lecturer_transition(self.new_panel, 'RETIRING')
        blocked = self.lecturer_transition(self.new_panel, 'RETIRED', 409)
        self.assertEqual(blocked.data['blockers']['pendingCoSupervisorNominations'], 1)
        self.post(self.office,
            f'/api/accounts/participants/lecturers/{self.new_panel.lecturer.staff_no}/pending-work/cancel/',
            {'recordType': 'CO_SUPERVISOR_NOMINATION', 'recordId': pending.pk,
             'reason': 'Resolve nomination before retirement'})
        self.lecturer_transition(self.new_panel, 'RETIRED')
        before = self.snapshot()
        self.post(self.new_panel, f'/api/appointments/co-supervisor/nominations/{pending.pk}/accept/', {}, 409)
        self.assertEqual(self.snapshot(), before)

    def test_deferred_student_direct_assignment_and_marks_writes_fail_closed(self):
        task = self.draft()
        self.transition('DEFERRED')
        before = self.snapshot()
        self.post(self.supervisor, '/api/appointments/co-supervisor/nominations/',
            {'studentId': self.student.pk, 'candidateId': self.new_panel.pk,
             'justification': 'Direct identifier bypass'}, 409)
        self.post(self.supervisor, '/api/appointments/panel/recommendations/',
            {'studentId': self.student.matric_no, 'recommendedMemberId': self.new_panel.lecturer.staff_no,
             'justification': 'Direct identifier bypass'}, 400)
        self.post(self.office,
            f'/api/appointments/supervisor/appointments/{self.supervisor_appointment.pk}/replacement/',
            {'proposedSupervisorId': self.new_panel.lecturer.staff_no, 'reason': 'Direct identifier bypass'}, 409)
        self.post(self.office, f'/api/marks/periods/{self.period.pk}/manual-overrides/',
            {'studentId': self.student.matric_no, 'evaluatorId': self.new_panel.pk,
             'reason': 'Direct identifier bypass'}, 400)
        self.post(self.office, '/api/marks/completion-windows/',
            {'taskIds': [task.pk], 'deadline': (timezone.now() + timedelta(days=10)).isoformat(),
             'reason': 'Direct identifier bypass'}, 409)
        self.post(self.panel, f'/api/marks/tasks/{task.pk}/submit/', submission_payload(task), 404)
        ensure_period_tasks(self.period, actor=self.office)
        self.assertEqual(self.snapshot(), before)

    def test_retiring_evaluator_cannot_be_selected_by_direct_backup_or_supporting_identifiers(self):
        self.lecturer_transition(self.new_panel, 'RETIRING')
        before = self.snapshot()
        self.post(self.office, f'/api/marks/periods/{self.period.pk}/manual-overrides/',
            {'studentId': self.student.matric_no, 'evaluatorId': self.new_panel.pk,
             'reason': 'Bypass hidden candidate'}, 400)
        self.post(self.supervisor, '/api/appointments/co-supervisor/nominations/',
            {'studentId': self.student.pk, 'candidateId': self.new_panel.pk,
             'justification': 'Bypass hidden candidate'}, 409)
        self.assertEqual(self.snapshot(), before)

    def test_terminal_history_is_readable_without_student_private_details_or_supporting_marks(self):
        task = self.draft()
        self.transition('WITHDRAWN')
        before = self.snapshot()
        dossier = f'/api/dashboard/progress/{self.student.matric_no}/'
        public = self.get(self.student_user, dossier).data
        self.assertEqual(public['student']['status'], 'Withdrawn')
        self.assertIsNone(public['marks'])
        self.assertNotIn('PRIVATE', str(public))
        self.get(self.new_supervisor, dossier, 404)
        own_history = self.get(self.new_supervisor, '/api/appointments/co-supervisor/').data
        self.assertEqual(own_history['appointments'][0]['status'], 'ENDED')
        self.get(self.new_supervisor, f'/api/marks/records/MRK-{task.pk:05d}/', 403)
        self.assertEqual(self.get(self.new_supervisor, '/api/marks/my-submitted-evaluations/').data, [])
        submitted = self.get(self.supervisor, '/api/marks/my-submitted-evaluations/').data
        self.assertEqual([row['id'] for row in submitted], [self.entry.task_id])
        self.assertEqual(self.get(self.panel, '/api/marks/my-evaluation-tasks/').data, [])
        self.get(self.office, f'/api/marks/records/MRK-{task.pk:05d}/')
        office = self.get(self.office, dossier).data
        self.assertIn('PRIVATE lifecycle decision', str(office))
        report = self.get(self.office, '/api/dashboard/reports/?semester=all').data
        # Operational reports exclude retired drafts; Office record detail retains them.
        self.assertEqual(report['marks']['total'], 1)
        workbook = load_workbook(BytesIO(self.get(self.office,
            '/api/dashboard/reports/export/?semester=all').content))
        self.assertEqual(workbook['Marks'].max_row, 2)
        self.get(self.student_user, '/api/dashboard/reports/export/', 403)
        self.assertEqual(self.snapshot(), before)

    def test_student_request_history_does_not_expose_private_terminal_lifecycle_reason(self):
        from appointments.serializers import SupervisorApplicationSerializer
        self.transition('WITHDRAWN')
        self.application.refresh_from_db()
        self.assertNotIn('PRIVATE lifecycle decision', str(SupervisorApplicationSerializer(self.application).data))
        for path in ('/api/appointments/supervisor/applications/',
                     f'/api/appointments/supervisor/applications/{self.application.pk}/',
                     '/api/notifications/'):
            with self.subTest(path=path):
                self.assertNotIn('PRIVATE lifecycle decision', str(self.get(self.student_user, path).data))

    def test_cancelled_primary_replacement_does_not_leak_private_lifecycle_reason(self):
        replacement = self.post(self.office,
            f'/api/appointments/supervisor/appointments/{self.supervisor_appointment.pk}/replacement/',
            {'proposedSupervisorId': self.new_panel.lecturer.staff_no, 'reason': 'Primary change requested'}, 201)
        self.transition('WITHDRAWN')
        staff_detail = self.get(self.office,
            f"/api/appointments/supervisor/applications/{replacement.data['id']}/").data
        self.assertIn('PRIVATE lifecycle decision', str(staff_detail))
        for path in ('/api/appointments/supervisor/applications/',
                     f"/api/appointments/supervisor/applications/{replacement.data['id']}/",
                     f'/api/dashboard/progress/{self.student.matric_no}/'):
            with self.subTest(path=path):
                self.assertNotIn('PRIVATE lifecycle decision', str(self.get(self.student_user, path).data))

    def test_retiring_lecturer_direct_primary_and_panel_assignment_is_rejected(self):
        self.lecturer_transition(self.new_panel, 'RETIRING')
        self.post(self.office,
            f'/api/appointments/panel/appointments/{self.panel_appointment.pk}/end/',
            {'outcome': 'OTHER', 'reason': 'Vacate Panel position for candidate eligibility check'})
        before = self.snapshot()
        self.post(self.office,
            f'/api/appointments/supervisor/appointments/{self.supervisor_appointment.pk}/replacement/',
            {'proposedSupervisorId': self.new_panel.lecturer.staff_no, 'reason': 'Bypass candidate list'}, 409)
        self.post(self.supervisor, '/api/appointments/panel/recommendations/',
            {'studentId': self.student.matric_no, 'recommendedMemberId': self.new_panel.lecturer.staff_no,
             'justification': 'Bypass candidate list'}, 400)
        self.assertEqual(self.snapshot(), before)


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class CompoundLifecycleConcurrencyTests(CompoundLifecycleFixture, APITransactionTestCase):
    def compete(self, targets):
        barrier = Barrier(2)

        def transition(target):
            close_old_connections()
            try:
                with connection.cursor() as cursor:
                    cursor.execute("SET lock_timeout TO '5s'")
                actor = User.objects.get(pk=self.office.pk)
                barrier.wait(timeout=10)
                try:
                    transition_student(matric_no=self.student.matric_no, actor=actor,
                        target_status=target, reason=f'Concurrent {target}')
                    return ('changed', target)
                except ParticipantLifecycleConflict:
                    return ('conflict', target)
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(transition, target) for target in targets]
            results = [future.result(timeout=30) for future in futures]
        self.assertCountEqual([result[0] for result in results], ['changed', 'conflict'])
        return next(target for result, target in results if result == 'changed')

    def test_competing_terminal_transitions_have_one_audited_outcome(self):
        before = self.marks_snapshot()
        winner = self.compete(['GRADUATED', 'WITHDRAWN'])
        self.student.refresh_from_db()
        self.assertEqual(self.student.status.upper(), winner)
        self.assertEqual(ParticipantLifecycleAudit.objects.filter(student=self.student).count(), 1)
        self.assert_closed('COMPLETED' if winner == 'GRADUATED' else 'WITHDRAWN')
        self.assertEqual(self.marks_snapshot(), before)

    def test_marks_submission_competing_with_withdrawal_preserves_one_valid_outcome(self):
        task = self.draft()
        payload = submission_payload(task)
        barrier = Barrier(2)

        def run(submitting):
            close_old_connections()
            try:
                with connection.cursor() as cursor:
                    cursor.execute("SET lock_timeout TO '5s'")
                client = APIClient()
                client.force_authenticate(User.objects.get(pk=(self.panel if submitting else self.office).pk))
                barrier.wait(timeout=10)
                if submitting:
                    return client.post(f'/api/marks/tasks/{task.pk}/submit/', payload, format='json').status_code
                return client.post(
                    f'/api/accounts/participants/students/{self.student.matric_no}/transition/',
                    {'targetStatus': 'WITHDRAWN', 'reason': 'Concurrent withdrawal'}, format='json').status_code
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            submitting = pool.submit(run, True)
            withdrawing = pool.submit(run, False)
            submit_status, withdraw_status = submitting.result(timeout=30), withdrawing.result(timeout=30)
        self.assertEqual(withdraw_status, 200)
        self.assertIn(submit_status, (200, 404, 409))
        self.student.refresh_from_db()
        task.refresh_from_db()
        entry = MarkEntry.objects.get(task=task)
        self.assertEqual(self.student.status, Student.Status.WITHDRAWN)
        self.assert_closed('WITHDRAWN')
        self.assertEqual(entry.comments, 'Saved draft survives')
        self.assertEqual(entry.total_mark, 46)
        self.assertEqual(entry.scores.get().marks_awarded, 46)
        self.assertEqual((entry.status, task.lifecycle_status),
            ('SUBMITTED', 'ACTIVE') if submit_status == 200 else ('DRAFT', 'RETIRED'))
        self.assertEqual(MarkSubmissionReceipt.objects.filter(task=task).count(), int(submit_status == 200))
        self.assertFalse(task_access(task)['canEdit'])

    def test_duplicate_withdrawals_close_pending_support_and_marks_exactly_once(self):
        pending = self.pending_support()
        task = self.draft()
        self.window(task)
        before = self.marks_snapshot()
        self.compete(['WITHDRAWN', 'WITHDRAWN'])
        self.assert_closed('WITHDRAWN')
        pending.refresh_from_db()
        self.assertEqual(pending.status, CoSupervisorNomination.Status.CANCELLED)
        self.assertEqual(pending.workflow_events.filter(action='OFFICE_CANCEL_PARTICIPANT_LIFECYCLE').count(), 1)
        self.assertEqual(task.lifecycle_audits.filter(action='RETIRED').count(), 1)
        self.assertEqual(ParticipantLifecycleAudit.objects.filter(student=self.student).count(), 1)
        self.assertEqual(self.marks_snapshot(), before)

    def test_duplicate_deferrals_pause_and_invalidate_window_exactly_once(self):
        task = self.draft()
        window = self.window(task)
        before = self.marks_snapshot()
        self.compete(['DEFERRED', 'DEFERRED'])
        task.refresh_from_db()
        window.refresh_from_db()
        self.assertEqual(task.lifecycle_status, 'PAUSED')
        self.assertEqual(window.end_kind, 'INVALIDATED')
        self.assertEqual(task.lifecycle_audits.filter(action='PAUSED').count(), 1)
        self.assertEqual(ParticipantLifecycleAudit.objects.filter(student=self.student).count(), 1)
        self.assertEqual(self.marks_snapshot(), before)
