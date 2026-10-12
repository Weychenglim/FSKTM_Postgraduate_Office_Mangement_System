"""The Excel template for the Student Registry bulk import (UC04)."""

import io
import json

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from .models import Student, User
from .programmes import APPROVED_PROGRAMMES
from .registry_import import IMPORT_HEADERS
from .registry_template import XLSX_CONTENT_TYPE

DS, CS, _ = APPROVED_PROGRAMMES
TEMPLATE = "/api/registry/students/import/template/"
INSPECT = "/api/registry/students/import/inspect/"
IMPORT = "/api/registry/students/import/"


class ImportTemplateTests(APITestCase):
    def setUp(self):
        cache.clear()
        self.office = User.objects.create_user(
            email="template-office@example.test",
            password="pw-office-7901",
            full_name="Template Office",
            role=User.Role.OFFICE_ADMIN,
        )
        self.client.force_authenticate(self.office)

    def workbook(self):
        response = self.client.get(TEMPLATE)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return load_workbook(io.BytesIO(response.content))

    def test_the_template_opens_with_the_import_columns(self):
        response = self.client.get(TEMPLATE)
        self.assertEqual(response["Content-Type"], XLSX_CONTENT_TYPE)
        self.assertIn('filename="student-registry-template.xlsx"', response["Content-Disposition"])
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Students", "How to fill in", "Programmes"])
        students = book["Students"]
        self.assertEqual([cell.value for cell in students[1]], IMPORT_HEADERS)
        self.assertEqual(students.freeze_panes, "A2")
        self.assertEqual((students["A2"].number_format, students["E1001"].number_format), ("@", "@"))
        self.assertEqual(book["Programmes"].sheet_state, "hidden")

    def test_the_programme_column_is_a_dropdown_of_the_approved_programmes(self):
        book = self.workbook()
        validations = book["Students"].data_validations.dataValidation
        self.assertEqual(len(validations), 1)
        self.assertEqual((validations[0].type, validations[0].formula1), ("list", "=Programmes!$A$1:$A$3"))
        self.assertEqual(str(validations[0].sqref), "C2:C1001")
        self.assertEqual([row[0].value for row in book["Programmes"].iter_rows()], APPROVED_PROGRAMMES)

    def test_a_filled_in_template_imports_even_when_saved_on_the_help_sheet(self):
        book = self.workbook()
        students = book["Students"]
        rows = [
            ["WQD330001", "Template One", DS, "template.one@example.test", "0123456789", None, None],
            ["WQD330002", "Template Two", CS, None, None, "2026/2027 Semester I", None],
        ]
        for number, values in enumerate(rows, start=2):
            for column, value in enumerate(values, start=1):
                students.cell(row=number, column=column, value=value)
        book.active = book.sheetnames.index("How to fill in")
        buffer = io.BytesIO()
        book.save(buffer)

        def upload():
            return SimpleUploadedFile("smoke-template.xlsx", buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)

        layout = self.client.post(INSPECT, {"file": upload()}, format="multipart").data
        self.assertEqual((layout["sheet"], layout["layoutRecognised"], layout["rowCount"]), ("Students", True, 2))
        response = self.client.post(IMPORT, {"file": upload(), "options": json.dumps({})}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(Student.objects.get(matric_no="WQD330001").user.phone, "0123456789")
        self.assertEqual(Student.objects.get(matric_no="WQD330002").user.email, "wqd330002@siswa.um.edu.my")

    def test_only_the_office_downloads_it(self):
        lecturer = User.objects.create_user(
            email="template-lecturer@example.test",
            password="pw-lect-7902",
            full_name="Template Lecturer",
            role=User.Role.LECTURER,
        )
        self.client.force_authenticate(lecturer)
        self.assertEqual(self.client.get(TEMPLATE).status_code, status.HTTP_403_FORBIDDEN)
