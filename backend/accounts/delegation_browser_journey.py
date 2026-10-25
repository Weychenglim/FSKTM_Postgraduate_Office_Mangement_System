"""Opt-in real HTTP acting-delegation journeys; disposable databases only."""

from datetime import timedelta
from io import BytesIO
import json
from pathlib import Path
import subprocess
from unittest.mock import patch
from urllib.request import Request, urlopen

from django.db import connection
from django.test import LiveServerTestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework_simplejwt.tokens import AccessToken

from appointments import co_supervision
from appointments import test_appointment_lifecycle as primary_fixtures
from appointments.models import (
    CoSupervisorAppointment, CoSupervisorNomination, PanelRecommendation,
    StudentResearchProfile, SupervisorApplication, SupervisorAppointment,
)
from appointments.test_co_supervision_lifecycle import submitted_marks
from marks.models import (
    EvaluationTask, EvaluationTaskHandoverAudit, EvaluationTaskLifecycleAudit,
    MarkEntry, MarkScore,
)
from .delegations import grant_delegation
from .models import CoordinatorDelegation, Student, User


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class DelegationBrowserJourney(LiveServerTestCase):
    host = "127.0.0.1"

    def _lecturer(self, email, staff_no, *, supervisor=False, panel=False):
        return primary_fixtures.AppointmentLifecycleTests._lecturer(
            self, email, staff_no, supervisor=supervisor or panel, panel=panel
        )

    def setUp(self):
        self.assertTrue(
            str(connection.settings_dict["NAME"]).startswith("test_"),
            "A disposable test database is required.",
        )
        primary_fixtures.AppointmentLifecycleTests.setUp(self)
        self.office.is_staff = True
        self.office.save(update_fields=["is_staff"])
        self.acting = self.other_coordinator
        self.regular_programme = self.acting.lecturer.coordinator.programme_managed
        self.target_programme = self.student.programme
        self.today = timezone.localdate()
        self.pending = []
        for candidate in (self.new_supervisor,):
            row = co_supervision.nominate(
                actor=self.supervisor, student_id=self.student.pk,
                candidate_id=candidate.pk, justification="Delegation acceptance nomination",
            )
            co_supervision.decide(nomination_id=row.pk, actor=candidate, action="accept")
            self.pending.append(row.pk)
        self.target_request = self.pending_primary("DELEGATED-PENDING", self.target_programme)
        self.regular_request = self.pending_primary("REGULAR-PENDING", self.regular_programme)
        self.foreign_request = self.pending_primary(
            "FOREIGN-PENDING", "MASTER OF CYBER SECURITY (COURSEWORK)"
        )
        self.panel_record = PanelRecommendation.objects.create(
            profile=self.profile, academic_semester=self.semester,
            supervisor=self.supervisor, recommended_member=self.panel,
            status="PENDING_COORDINATOR", submitted_at=timezone.now(),
            panel_decided_at=timezone.now(),
        )
        self.entry, self.period = submitted_marks(self)
        self.preserved_models = (
            SupervisorApplication, SupervisorAppointment, PanelRecommendation,
            StudentResearchProfile, EvaluationTask, MarkEntry, MarkScore,
            EvaluationTaskLifecycleAudit, EvaluationTaskHandoverAudit,
        )
        self.before = {
            model: list(model.objects.order_by("pk").values())
            for model in self.preserved_models
        }

    def pending_primary(self, matric, programme):
        user = User.objects.create_user(
            email=f"{matric.lower()}@delegation-journey.example.test",
            full_name=matric, role=User.Role.STUDENT,
        )
        student = Student.objects.create(user=user, matric_no=matric, programme=programme)
        return SupervisorApplication.objects.create(
            student=student, academic_semester=self.semester,
            proposed_supervisor=self.new_supervisor,
            research_title="Pending delegated supervision", research_area="Software",
            research_abstract="Fictional pending primary request.",
            status="PENDING_COORDINATOR", supervisor_decided_at=timezone.now(),
        )

    def run_journey(self, width, mode):
        actors = {
            "office": self.office, "acting": self.acting,
            "regular": self.coordinator, "student": self.student_user,
            "primary": self.supervisor, "second": self.new_panel,
        }
        config = {
            "backend": self.live_server_url, "width": width, "mode": mode,
            "tokens": {key: str(AccessToken.for_user(user)) for key, user in actors.items()},
            "actingId": self.acting.pk, "actingName": self.acting.full_name,
            "targetProgramme": self.target_programme,
            "regularProgramme": self.regular_programme,
            "studentId": self.student.pk, "matricNo": self.student.matric_no,
            "targetMatric": "DELEGATED-PENDING", "regularMatric": "REGULAR-PENDING",
            "foreignMatric": "FOREIGN-PENDING",
            "targetRequestId": self.target_request.pk,
            "regularRequestId": self.regular_request.pk,
            "supportingIds": self.pending,
            "secondCandidateId": self.new_panel.pk,
            "today": self.today.isoformat(),
            "endsOn": (self.today + timedelta(days=2)).isoformat(),
        }
        completed = subprocess.run(
            ["node", "tests/delegationJourney.browser.mjs"],
            cwd=Path(__file__).resolve().parents[2] / "frontend",
            input=json.dumps(config), text=True, capture_output=True, timeout=300,
        )
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        print(completed.stdout.strip())
        for model, snapshot in self.before.items():
            self.assertEqual(list(model.objects.order_by("pk").values()), snapshot, model.__name__)
        # Parse a real authenticated XLSX response, not just its content type.
        token = str(AccessToken.for_user(self.acting))
        response = urlopen(Request(
            self.live_server_url + "/api/dashboard/reports/export/",
            headers={"Authorization": "Bearer " + token},
        ), timeout=30)
        with response:
            workbook = load_workbook(BytesIO(response.read()))
        ids = {str(cell.value) for cell in workbook["Supervisor"]["A"][1:]}
        self.assertIn(str(self.regular_request.pk), ids)
        self.assertNotIn(str(self.foreign_request.pk), ids)
        if mode == "ACTIVE":
            self.assertIn(str(self.target_request.pk), ids)
        else:
            self.assertNotIn(str(self.target_request.pk), ids)
        self.assertEqual(
            workbook.sheetnames, ["Summary", "Research Amendments", "Supervisor", "Panel"]
        )

    def revocation(self, width):
        self.run_journey(width, "revocation")
        grants = list(CoordinatorDelegation.objects.order_by("pk"))
        self.assertEqual(len(grants), 3)
        for grant in grants:
            self.assertEqual(grant.revoked_by_id, self.office.pk)
            self.assertEqual(grant.granted_by_id, self.office.pk)
            self.assertEqual(grant.revocation_reason, "Delegation acceptance revocation")
        first, second = list(CoSupervisorNomination.objects.order_by("pk"))
        self.assertEqual((first.status, second.status), ("APPROVED", "PENDING_COORDINATOR"))
        self.assertEqual((first.workflow_events.count(), second.workflow_events.count()), (3, 2))
        self.assertEqual(first.workflow_events.get(action="APPROVE").actor_id, self.acting.pk)
        appointment = CoSupervisorAppointment.objects.get(nomination=first)
        self.assertEqual(appointment.approved_by_id, self.acting.pk)
        self.assertEqual(appointment.status, "ACTIVE")

    def expiry(self, width):
        grant = grant_delegation(
            actor=self.office, coordinator_id=self.acting.pk,
            programme=self.target_programme, starts_on=self.today + timedelta(days=1),
            ends_on=self.today + timedelta(days=2), justification="Scheduled expiry acceptance",
        )
        before = list(CoordinatorDelegation.objects.values())
        for day, mode in (
            (self.today, "SCHEDULED"),
            (grant.ends_on, "ACTIVE"),
            (grant.ends_on + timedelta(days=1), "EXPIRED"),
        ):
            with patch("accounts.delegations.today", return_value=day):
                self.run_journey(width, mode)
        self.assertEqual(list(CoordinatorDelegation.objects.values()), before)
        for pk in self.pending:
            row = CoSupervisorNomination.objects.get(pk=pk)
            self.assertEqual(row.status, "PENDING_COORDINATOR")
            self.assertEqual(row.workflow_events.count(), 2)
        self.assertFalse(CoSupervisorAppointment.objects.exists())

    def test_revocation_desktop(self):
        self.revocation(1440)

    def test_revocation_phone(self):
        self.revocation(390)

    def test_scheduled_inclusive_expiry_desktop(self):
        self.expiry(1440)

    def test_scheduled_inclusive_expiry_phone(self):
        self.expiry(390)
