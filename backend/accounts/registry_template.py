"""The Excel template for the Student Registry bulk import (UC04).

The office can import its own spreadsheet as it is; the template is for when a
list has to be typed. Student ID and Phone are text columns, so Excel keeps
their leading zeros, and Programme is a dropdown of the approved programmes.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .programmes import APPROVED_PROGRAMMES
from .registry_import import IMPORT_HEADERS, MAX_IMPORT_ROWS, student_email_domain

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEMPLATE_FILE_NAME = "student-registry-template.xlsx"
TEXT_COLUMNS = ("student_id", "phone")
WIDTHS = {
    "student_id": 16,
    "full_name": 32,
    "programme": 46,
    "email": 34,
    "phone": 16,
    "intake": 24,
    "current_semester": 30,
}


def _guide_rows():
    # One column name per row, so this sheet is never read as the student list.
    domain = student_email_domain()
    return [
        ("Column", "What to enter", "Example"),
        ("student_id", "Required. The Student ID, in letters and digits (no spaces or /).", "23004955"),
        ("full_name", "Required. The student's full name.", "Aisyah binti Rahman"),
        ("programme", "Required. Pick one from the list in the cell.", APPROVED_PROGRAMMES[0]),
        (
            "email",
            f"Leave it blank to use the student's SiswaMail address (Student ID@{domain}).",
            f"23004955@{domain}",
        ),
        ("phone", "Optional.", "012-3456789"),
        ("intake", "Optional.", "2026/2027 Semester I"),
        ("current_semester", "Optional.", "Semester I, 2026/2027 Session"),
    ]


def build_template():
    """The template as ``.xlsx`` bytes: Students, How to fill in, and a hidden
    Programmes sheet that feeds the dropdown."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Students"
    sheet.append(IMPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E293B")
    sheet.freeze_panes = "A2"
    last_row = MAX_IMPORT_ROWS + 1
    for number, name in enumerate(IMPORT_HEADERS, start=1):
        letter = get_column_letter(number)
        sheet.column_dimensions[letter].width = WIDTHS[name]
        if name in TEXT_COLUMNS:
            for row in range(2, last_row + 1):
                sheet[f"{letter}{row}"].number_format = "@"

    guide = book.create_sheet("How to fill in")
    guide.append(["Fill in the Students sheet, one student per row, then upload the file. Keep its first row as it is."])
    guide.append([])
    for row in _guide_rows():
        guide.append(list(row))
    for cell in guide[3]:
        cell.font = Font(bold=True)
    for letter, width in (("A", 20), ("B", 80), ("C", 40)):
        guide.column_dimensions[letter].width = width
    for row in guide.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    choices = book.create_sheet("Programmes")
    for programme in APPROVED_PROGRAMMES:
        choices.append([programme])
    choices.sheet_state = "hidden"
    column = get_column_letter(IMPORT_HEADERS.index("programme") + 1)
    validation = DataValidation(
        type="list",
        formula1=f"=Programmes!$A$1:$A${len(APPROVED_PROGRAMMES)}",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Programme",
        error="Pick a programme from the list.",
        showInputMessage=True,
        promptTitle="Programme",
        prompt="Pick from the list.",
    )
    sheet.add_data_validation(validation)
    validation.add(f"{column}2:{column}{last_row}")

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
