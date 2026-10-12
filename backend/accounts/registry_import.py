"""Student Registry bulk import (UC04): read a CSV or XLSX file and check rows.

The file may be our template or the office's own extract: column names, order,
extra columns, title rows above the column names and several sheets are all
handled, and the office confirms which column is which (``options``).

Rows are never auto-corrected. A missing email or an unrecognised programme is
reported so a person fixes it; inventing a value would attach a real account to
the wrong mailbox.
"""

import csv
import datetime
import io
import json
import os
import re
import zipfile

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db.models.functions import Lower
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import Lecturer, OfficeStaff, Student
from .programmes import normalise_programme
from .registry_columns import (
    FIELD_KEYS,
    FIELD_LABELS,
    FIELDS,
    column_letter,
    field_for_header,
    header_score,
    suggest_mapping,
)
from .validators import matric_problem, phone_problem

IMPORT_HEADERS = ["student_id", "full_name", "programme", "email", "phone", "intake", "current_semester"]
# The original template had only these; files without the last two still import.
REQUIRED_TEMPLATE_HEADERS = IMPORT_HEADERS[:5]
ALLOWED_EXTENSIONS = {".csv", ".xlsx"}
DEFAULT_MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
# Same limits as single registration, so the preview predicts the commit.
MAX_MATRIC_LENGTH = 64
MAX_NAME_LENGTH = 255
MAX_EMAIL_LENGTH = 254
MAX_PHONE_LENGTH = 32
MAX_INTAKE_LENGTH = 32
MAX_SEMESTER_LENGTH = 64
# An XLSX is a ZIP: cap what it may expand to, how wide a row is read, and how
# many blank rows are walked, so a tiny padded or bomb workbook cannot pin a worker.
MAX_XLSX_UNCOMPRESSED_BYTES = 30 * 1024 * 1024
MAX_XLSX_COLUMNS = 50
MAX_BLANK_RUN = 200
# The column names must be among the first non-blank rows (title rows above
# them are skipped).
HEADER_SCAN_ROWS = 15
CSV_DELIMITERS = (",", ";", "\t")
SAMPLE_COUNT = 3
# Settings the office can give for every row: option name -> field it fills.
DEFAULT_FIELDS = {"programme": "programme", "intake": "intake", "semester": "current_semester"}
UTF8_MULTIBYTE = re.compile(rb"[\xc2-\xf4][\x80-\xbf]+")


class ImportFileError(ValueError):
    pass


class ImportOptionsError(ValueError):
    """Bad import options, held as ``{field: [message]}`` for a 400 response."""

    def __init__(self, errors):
        super().__init__(" ".join(message for messages in errors.values() for message in messages))
        self.errors = errors


def max_import_bytes():
    return getattr(settings, "REGISTRY_IMPORT_MAX_BYTES", DEFAULT_MAX_IMPORT_BYTES)


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.date().isoformat() if isinstance(value, datetime.datetime) else value.isoformat()
    return str(value).strip()


def _decode_csv(raw):
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        # Excel's "CSV (Comma delimited)" saves Windows-1252, which is fine. A
        # UTF-8 file with one stray byte is not: decoding all of it as
        # Windows-1252 would silently turn "José" into "JosÃ©".
        if UTF8_MULTIBYTE.search(raw):
            line = raw[: exc.start].count(b"\n") + 1
            raise ImportFileError(
                f"Line {line} is not valid UTF-8. Save the file as \"CSV UTF-8\" and try again."
            ) from exc
    try:
        return raw.decode("cp1252")
    except UnicodeDecodeError as exc:
        raise ImportFileError("The CSV file could not be read. Save it as UTF-8 and try again.") from exc


def _csv_delimiter(text):
    """Comma, semicolon or tab: whichever splits the first lines into the most
    cells. Cells pasted from a spreadsheet are tab-separated, and some regional
    Excel settings save CSV with semicolons."""
    lines = [line for line in text.splitlines()[: HEADER_SCAN_ROWS * 2] if line.strip()]
    best, best_width = ",", 1
    for delimiter in CSV_DELIMITERS:
        width = 0
        for line in lines[:HEADER_SCAN_ROWS]:
            try:
                width = max(width, len(next(csv.reader([line], delimiter=delimiter))))
            except csv.Error:
                continue
        if width > best_width:
            best, best_width = delimiter, width
    return best


def _csv_lines(raw):
    text = _decode_csv(raw)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_csv_delimiter(text))
    try:
        for cells in reader:
            yield reader.line_num, [cell.strip() for cell in cells]
    except csv.Error as exc:
        raise ImportFileError(
            f"Line {reader.line_num + 1} of the CSV could not be read ({exc}). "
            "Check its quotes and save it as CSV UTF-8."
        ) from exc


def _open_workbook(raw):
    if not raw.startswith(b"PK\x03\x04"):
        raise ImportFileError("That file is not a valid XLSX workbook.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            expanded = sum(entry.file_size for entry in archive.infolist())
    except (zipfile.BadZipFile, OSError) as exc:
        raise ImportFileError("That file is not a valid XLSX workbook.") from exc
    if expanded > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ImportFileError("That workbook is too large to import. Save it as CSV and try again.")
    try:
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ImportFileError("That file is not a valid XLSX workbook.") from exc


def _sheet_lines(sheet, limit):
    """A sheet's non-blank rows as ``(row number, cells)``, stopping after
    ``limit`` rows or a long blank run so a padded sheet cannot pin a worker."""
    kept = blank_run = 0
    rows = sheet.iter_rows(values_only=True, max_col=MAX_XLSX_COLUMNS)
    for line, values in enumerate(rows, start=1):
        cells = [_text(value) for value in values]
        if not any(cells):
            blank_run += 1
            if blank_run > MAX_BLANK_RUN:
                return
            continue
        blank_run = 0
        yield line, cells
        kept += 1
        if kept >= limit:
            return


def _pick_sheet(workbook, requested):
    """The requested sheet, else the first sheet (the open one first) whose top
    rows name at least two import fields, else the open sheet."""
    names = workbook.sheetnames
    if requested is not None:
        if requested not in names:
            raise ImportOptionsError({"sheet": [f'This workbook has no sheet called "{requested}".']})
        return requested
    active = workbook.active.title
    for name in [active, *(name for name in names if name != active)]:
        if any(header_score(cells) >= 2 for _, cells in _sheet_lines(workbook[name], HEADER_SCAN_ROWS)):
            return name
    return active


def _check_upload(uploaded_file):
    extension = os.path.splitext(uploaded_file.name or "")[1].lower()
    if extension == ".xls":
        raise ImportFileError(
            "This is an older Excel file (.xls). Open it in Excel, save it as .xlsx and import that file."
        )
    if extension not in ALLOWED_EXTENSIONS:
        raise ImportFileError("Only XLSX or CSV files allowed.")
    limit = max_import_bytes()
    if uploaded_file.size > limit:
        raise ImportFileError(f"The file is larger than the {limit // 1024} KB import limit.")
    return extension


def load_import_lines(uploaded_file, sheet=None):
    """``(sheet names, chosen sheet, [(row number, cells)])`` for the file's
    non-blank rows. A CSV has no sheets."""
    extension = _check_upload(uploaded_file)
    raw = uploaded_file.read()
    if extension == ".csv":
        sheets, chosen = [], None
        lines = [(line, cells) for line, cells in _csv_lines(raw) if any(cells)]
    else:
        workbook = _open_workbook(raw)
        try:
            sheets = list(workbook.sheetnames)
            chosen = _pick_sheet(workbook, sheet)
            # Enough rows to see past title rows and still notice a file over the limit.
            lines = list(_sheet_lines(workbook[chosen], MAX_IMPORT_ROWS + HEADER_SCAN_ROWS + 1))
        finally:
            workbook.close()
    if not lines:
        raise ImportFileError("The file is empty.")
    return sheets, chosen, lines


def detect_header_index(lines):
    """The column-name row among the first non-blank rows: the one naming the
    most import fields (at least two), otherwise the fullest of the first ten."""
    top = lines[:HEADER_SCAN_ROWS]
    scores = [header_score(cells) for _, cells in top]
    if max(scores) >= 2:
        return scores.index(max(scores))
    filled = [sum(1 for cell in cells if cell) for _, cells in top[:10]]
    return filled.index(max(filled))


def _header_index(lines, header_row):
    if header_row is None:
        return detect_header_index(lines)
    for index, (line, _) in enumerate(lines[:HEADER_SCAN_ROWS]):
        if line == header_row:
            return index
    raise ImportOptionsError(
        {"headerRow": [f"Row {header_row} is empty or too far down. Choose one of the first rows."]}
    )


def _width(header, rows):
    """How many columns hold anything, in the header or any data row."""
    def used(cells):
        return max((index + 1 for index, cell in enumerate(cells) if cell), default=0)

    return min(MAX_XLSX_COLUMNS, max([used(header), *(used(cells) for _, cells in rows)]))


def _student_rows(lines, header_index):
    rows = lines[header_index + 1:]
    if not rows:
        raise ImportFileError("The file has a header row but no student rows.")
    if len(rows) > MAX_IMPORT_ROWS:
        raise ImportFileError(
            f"The file has more than {MAX_IMPORT_ROWS} student rows. Split it and import each part."
        )
    return rows


def check_mapping(mapping, width, required):
    """Refuse column matches that point past the file, reuse a column, or leave
    a required field without a column."""
    problems, used = [], {}
    for field in FIELD_KEYS:
        index = mapping.get(field)
        if index is None:
            continue
        if index >= width:
            problems.append(f"{FIELD_LABELS[field]}: column {column_letter(index)} is not in the file.")
        elif index in used:
            problems.append(
                f"Column {column_letter(index)} is chosen for both {FIELD_LABELS[used[index]]} and {FIELD_LABELS[field]}."
            )
        else:
            used[index] = field
    missing = [FIELD_LABELS[field] for field in required if mapping.get(field) is None]
    if missing:
        problems.append(f"Choose a column for: {', '.join(missing)}.")
    if problems:
        raise ImportOptionsError({"mapping": problems})


def required_fields(options):
    """Fields that need a column. A programme set for every row needs none."""
    defaults = (options or {}).get("defaults") or {}
    covered = {"programme"} if defaults.get("programme") else set()
    return [key for key, _, required in FIELDS if required and key not in covered]


def _apply_defaults(cells, defaults):
    """Fill blank cells (or a column the file does not have) from the settings
    for every row. A value in the file is never overwritten."""
    for option, field in DEFAULT_FIELDS.items():
        if defaults.get(option) and not cells[field]:
            cells[field] = defaults[option]
    return cells


def read_import_file(uploaded_file, options=None):
    """Return ``[(row number, {field: value})]`` for every student row.

    Without ``options`` the file must be our template (column names in the
    first row), as before. With ``options`` any layout works: the column names
    may sit below title rows, and ``options["mapping"]`` says which column holds
    which field (suggested from the names when it is left out).
    """
    if options is None:
        _, _, lines = load_import_lines(uploaded_file)
        header = [cell.lower() for cell in lines[0][1]]
        missing = [name for name in REQUIRED_TEMPLATE_HEADERS if name not in header]
        if missing:
            raise ImportFileError(
                f"Missing required column(s): {', '.join(missing)}. "
                "Download the template for the expected format."
            )
        mapping = {name: header.index(name) if name in header else None for name in IMPORT_HEADERS}
        header_index = 0
    else:
        _, _, lines = load_import_lines(uploaded_file, options.get("sheet"))
        header_index = _header_index(lines, options.get("headerRow"))
        header = lines[header_index][1]
        mapping = options.get("mapping") or suggest_mapping(header)
        check_mapping(mapping, _width(header, lines[header_index + 1:]), required_fields(options))

    defaults = (options or {}).get("defaults") or {}
    return [
        (
            line,
            _apply_defaults(
                {
                    field: cells[mapping[field]] if mapping.get(field) is not None and mapping[field] < len(cells) else ""
                    for field in IMPORT_HEADERS
                },
                defaults,
            ),
        )
        for line, cells in _student_rows(lines, header_index)
    ]


def inspect_import_file(uploaded_file, options=None):
    """Describe a file's layout so the office can confirm which column is which.

    Only column names and a few sample values per column are returned, never
    whole rows.
    """
    options = options or {}
    sheets, chosen, lines = load_import_lines(uploaded_file, options.get("sheet"))
    header_index = _header_index(lines, options.get("headerRow"))
    header_line, header = lines[header_index]
    rows = _student_rows(lines, header_index)
    suggested = suggest_mapping(header)
    columns = []
    for index in range(_width(header, rows)):
        name = header[index] if index < len(header) else ""
        samples = []
        for _, cells in rows:
            if index < len(cells) and cells[index]:
                samples.append(cells[index][:60])
                if len(samples) == SAMPLE_COUNT:
                    break
        columns.append(
            {
                "index": index,
                "letter": column_letter(index),
                "header": name,
                "samples": samples,
                "suggestedField": field_for_header(name),
            }
        )
    matched = {index for index in suggested.values() if index is not None}
    ignored = [
        column["header"] or f"Column {column['letter']}"
        for column in columns
        if column["index"] not in matched and (column["header"] or column["samples"])
    ]
    return {
        "fileName": uploaded_file.name,
        "sheets": sheets,
        "sheet": chosen,
        "headerRow": header_line,
        "headerRowOptions": [
            {"row": line, "preview": " | ".join(cell for cell in cells if cell)[:80]}
            for line, cells in lines[:HEADER_SCAN_ROWS]
        ],
        "columns": columns,
        "fields": [{"key": key, "label": label, "required": required} for key, label, required in FIELDS],
        "mapping": suggested,
        "ignoredColumns": ignored,
        "rowCount": len(rows),
        # Our own template: names in the first row, every column known, nothing
        # left over. The office can go straight to checking rows.
        "layoutRecognised": header_index == 0
        and not ignored
        and all(suggested[key] is not None for key in required_fields(options)),
    }


def parse_import_options(raw):
    """Read the ``options`` form field (JSON). None when it is absent."""
    if raw is None or raw == "":
        return None
    try:
        data = json.loads(raw) if isinstance(raw, (str, bytes)) else raw
    except ValueError:
        data = None
    if not isinstance(data, dict):
        raise ImportOptionsError({"options": ["The import options could not be read."]})

    errors = {}
    sheet = data.get("sheet")
    if sheet is not None and not isinstance(sheet, str):
        errors["sheet"] = ["Choose a sheet by its name."]
    header_row = data.get("headerRow")
    if header_row is not None and (isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 1):
        errors["headerRow"] = ["Choose the row that holds the column names."]
    mapping = data.get("mapping")
    if mapping is not None:
        if not isinstance(mapping, dict):
            errors["mapping"] = ["The column matches could not be read."]
        elif any(key not in FIELD_KEYS for key in mapping):
            unknown = ", ".join(sorted(key for key in mapping if key not in FIELD_KEYS))
            errors["mapping"] = [f"Unknown field: {unknown}."]
        elif any(
            value is not None and (isinstance(value, bool) or not isinstance(value, int) or value < 0)
            for value in mapping.values()
        ):
            errors["mapping"] = ["Each match must be a column number."]
    defaults = _parse_defaults(data.get("defaults"), errors)
    send_invites = data.get("sendInvites", True)
    if not isinstance(send_invites, bool):
        errors["sendInvites"] = ["Choose whether to send the activation emails now."]
    if errors:
        raise ImportOptionsError(errors)
    return {
        "sheet": sheet,
        "headerRow": header_row,
        "mapping": {key: mapping.get(key) for key in FIELD_KEYS} if mapping is not None else None,
        "defaults": defaults,
        "sendInvites": send_invites,
    }


def _parse_defaults(raw, errors):
    """Settings for every row: an approved programme, an intake and a current
    semester, each optional and checked like single registration."""
    if raw is None:
        return {}
    if not isinstance(raw, dict) or any(key not in DEFAULT_FIELDS for key in raw):
        errors["defaults"] = ["The settings for every row could not be read."]
        return {}
    values = {key: str(raw.get(key) or "").strip() for key in DEFAULT_FIELDS}
    problems = []
    programme = normalise_programme(values["programme"]) if values["programme"] else ""
    if values["programme"] and not programme:
        problems.append(f'"{values["programme"]}" is not an approved programme.')
    if len(values["intake"]) > MAX_INTAKE_LENGTH:
        problems.append(f"Intake is longer than {MAX_INTAKE_LENGTH} characters.")
    if len(values["semester"]) > MAX_SEMESTER_LENGTH:
        problems.append(f"Current semester is longer than {MAX_SEMESTER_LENGTH} characters.")
    if problems:
        errors["defaults"] = problems
    values["programme"] = programme
    return {key: value for key, value in values.items() if value}


def _field_problem(matric, name, phone, intake="", semester=""):
    """The first length/format problem in a row's free-text fields, or None."""
    if len(matric) > MAX_MATRIC_LENGTH:
        return f"student_id is longer than {MAX_MATRIC_LENGTH} characters"
    problem = matric_problem(matric)
    if problem:
        return problem
    if len(name) > MAX_NAME_LENGTH:
        return f"full_name is longer than {MAX_NAME_LENGTH} characters"
    if len(phone) > MAX_PHONE_LENGTH:
        return f"phone is longer than {MAX_PHONE_LENGTH} characters"
    problem = phone_problem(phone)
    if problem:
        return problem
    if len(intake) > MAX_INTAKE_LENGTH:
        return f"intake is longer than {MAX_INTAKE_LENGTH} characters"
    if len(semester) > MAX_SEMESTER_LENGTH:
        return f"current_semester is longer than {MAX_SEMESTER_LENGTH} characters"
    return None


def _email_problem(email):
    if len(email) > MAX_EMAIL_LENGTH:
        return f"email is longer than {MAX_EMAIL_LENGTH} characters"
    try:
        validate_email(email)
    except ValidationError:
        return f'"{email}" is not a valid email address'
    return None


def _same_text(first, second):
    return " ".join((first or "").split()).casefold() == " ".join((second or "").split()).casefold()


def _existing_record(user):
    """The registry record an already-registered row matched, as the office sees it."""
    student = user._related_or_none("student")
    return {
        "id": student.matric_no if student else "",
        "name": user.full_name,
        "email": user.email,
        "programme": student.programme if student else "",
        "academicStatus": student.status if student else "",
        "role": user.role,
    }


def _differences(cells, existing, matched_by):
    """Identity details in the file that disagree with the record it matched,
    which usually means a wrong Student ID or email in the file."""
    checks = [("name", cells["full_name"], existing["name"])]
    if matched_by == "studentId":
        checks.append(("email", cells["email"], existing["email"]))
    else:
        checks.append(("id", cells["student_id"], existing["id"]))
    return [field for field, in_file, on_record in checks if in_file and not _same_text(in_file, on_record)]


def validate_import_rows(rows):
    """Check each row against the same rules as single registration."""
    first_id, first_email = {}, {}
    for line, cells in rows:
        if cells["student_id"]:
            first_id.setdefault(cells["student_id"].lower(), line)
        if cells["email"]:
            first_email.setdefault(cells["email"].lower(), line)

    # The matched accounts are kept, not just their keys, so already-registered
    # rows can show the office who they are. One query each, whatever the size.
    registered_ids = {
        student.key: student.user
        for student in Student.objects.select_related("user")
        .annotate(key=Lower("matric_no"))
        .filter(key__in=list(first_id))
    }
    registered_emails = {
        user.key: user
        for user in get_user_model()
        .objects.select_related("student")
        .annotate(key=Lower("email"))
        .filter(key__in=list(first_email))
    }
    # Sign-in accepts a matric or a staff number, so one value cannot be both.
    staff_numbers = {
        number.lower()
        for model in (OfficeStaff, Lecturer)
        for number in model.objects.annotate(key=Lower("staff_no"))
        .filter(key__in=list(first_id))
        .values_list("staff_no", flat=True)
    }

    checked = []
    for line, cells in rows:
        matric, name, email = cells["student_id"], cells["full_name"], cells["email"]
        programme = normalise_programme(cells["programme"])
        status, issue = "Ready", ""
        existing_user = matched_by = duplicate_of = None

        intake, semester = cells.get("intake", ""), cells.get("current_semester", "")
        field_problem = (
            _field_problem(matric, name, cells["phone"], intake, semester) if matric and name else None
        )
        email_problem = _email_problem(email) if email else None
        # A student who is already registered is skipped whatever else the row
        # says, so that is reported before problems the office would otherwise
        # be asked to fix for nothing (an extract re-uploaded for late
        # registrations lists everyone again).
        if not matric or not name:
            status = "Missing Field"
            issue = "student_id is required" if not matric else "full_name is required"
        elif field_problem:
            status = "Invalid Field"
            issue = field_problem
        elif first_id[matric.lower()] != line:
            status = "Duplicate In File"
            duplicate_of = first_id[matric.lower()]
            issue = f"same student_id as line {duplicate_of}"
        elif matric.lower() in staff_numbers:
            status = "Invalid Field"
            issue = "this student_id is already used as a staff number"
        elif matric.lower() in registered_ids:
            status = "Already Registered"
            issue = "this matric number is already registered"
            existing_user, matched_by = registered_ids[matric.lower()], "studentId"
        elif not email:
            status = "Missing Email"
            issue = "email is required; the student needs it to activate the account"
        elif email_problem:
            status = "Missing Email"
            issue = email_problem
        elif first_email[email.lower()] != line:
            status = "Duplicate In File"
            duplicate_of = first_email[email.lower()]
            issue = f"same email as line {duplicate_of}"
        elif email.lower() in registered_emails:
            status = "Already Registered"
            issue = "this email already has an account"
            existing_user, matched_by = registered_emails[email.lower()], "email"
        elif not programme:
            status = "Bad Programme"
            issue = f'"{cells["programme"]}" is not an approved programme'

        row = {
            "line": line,
            "id": matric,
            "name": name,
            "programme": programme or cells["programme"],
            "email": email,
            "phone": cells["phone"],
            "intake": intake,
            "semester": semester,
            "status": status,
            "issue": issue,
        }
        if duplicate_of is not None:
            row["duplicateOf"] = duplicate_of
        if existing_user is not None:
            row["existing"] = _existing_record(existing_user)
            row["matchedBy"] = matched_by
            row["differences"] = _differences(cells, row["existing"], matched_by)
        checked.append(row)
    return checked
