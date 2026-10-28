"""The review list for the office: every entry that needs checking, with the
PDF's own text and empty columns for the corrections."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import FaqEntry

COLUMNS = [
    ("Entry number", 13),
    ("Category", 24),
    ("Status", 15),
    ("Question now", 45),
    ("Answer now", 60),
    ("Text in the PDF", 60),
    ("What looks wrong", 35),
    ("Correct question (only if wrong)", 40),
    ("Correct answer", 60),
]
SAFE_START = ("=", "+", "-", "@")


def _cell(value):
    # Text from the PDF must never run as a spreadsheet formula.
    text = str(value or "")
    return f"'{text}" if text.startswith(SAFE_START) else text


def review_workbook(include_all=False):
    entries = FaqEntry.objects.select_related("category").exclude(status=FaqEntry.Status.ARCHIVED)
    if not include_all:
        entries = entries.filter(status=FaqEntry.Status.NEEDS_CHECKING)
    book = Workbook()
    sheet = book.active
    sheet.title = "FAQ to check"
    sheet.append([name for name, _ in COLUMNS])
    for index, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
        heading = sheet.cell(row=1, column=index)
        heading.font = Font(bold=True)
        heading.fill = PatternFill("solid", fgColor="E2E8F0")
    for entry in entries:
        sheet.append(
            [
                entry.pk,
                entry.category.name,
                entry.get_status_display(),
                _cell(entry.question),
                _cell(entry.answer),
                _cell(entry.source_text),
                _cell(entry.source_note),
                "",
                "",
            ]
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    output = BytesIO()
    book.save(output)
    return output.getvalue(), entries.count()
