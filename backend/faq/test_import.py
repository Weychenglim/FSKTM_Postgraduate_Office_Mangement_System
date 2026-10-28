"""Importing the office's FAQ (C1): statuses, categories, repeats and running twice."""
import json
import tempfile
from io import BytesIO, StringIO
from pathlib import Path

from django.core.management import call_command
from django.test import TestCase
from openpyxl import load_workbook

from .importer import CATEGORIES, ImportProblem, default_path, import_faq
from .models import FaqCategory, FaqChange, FaqEntry
from .review_export import review_workbook

Status = FaqEntry.Status


class ImportFaqTests(TestCase):
    def test_counts_per_status_match_the_plan(self):
        counts = import_faq()
        self.assertEqual(counts, {"published": 56, "needsChecking": 61, "repeatsSkipped": 46})
        self.assertEqual(FaqEntry.objects.count(), 117)
        self.assertEqual(FaqEntry.objects.filter(status=Status.PUBLISHED).count(), 56)
        self.assertEqual(FaqEntry.objects.filter(status=Status.NEEDS_CHECKING).count(), 61)
        self.assertEqual(FaqChange.objects.filter(action=FaqChange.Action.IMPORTED).count(), 117)

    def test_categories_have_plain_names_in_pdf_order(self):
        import_faq()
        self.assertEqual(list(FaqCategory.objects.values_list("name", flat=True)), list(CATEGORIES.values()))
        self.assertEqual(
            FaqEntry.objects.filter(source_ref__contains="categorized", category__name="Student application").count(),
            28,
        )

    def test_running_twice_adds_nothing_and_keeps_the_office_edits(self):
        import_faq()
        entry = FaqEntry.objects.get(source_ref="student-application-categorized-1")
        entry.answer = "Corrected by the office."
        entry.save()
        counts = import_faq()
        self.assertEqual(counts, {"alreadyThere": 117, "repeatsSkipped": 46})
        self.assertEqual(FaqEntry.objects.count(), 117)
        entry.refresh_from_db()
        self.assertEqual(entry.answer, "Corrected by the office.")

    def test_needs_checking_entries_are_never_searchable(self):
        import_faq()
        published = FaqEntry.objects.published()
        self.assertEqual(published.count(), 56)
        self.assertFalse(published.exclude(status=Status.PUBLISHED).exists())
        damaged = FaqEntry.objects.get(source_ref="student-application-categorized-10")
        self.assertEqual(damaged.status, Status.NEEDS_CHECKING)
        self.assertNotIn(damaged, published)

    def test_damaged_rows_keep_the_pdf_text_and_the_problem(self):
        import_faq()
        damaged = FaqEntry.objects.get(source_ref="student-application-categorized-10")
        self.assertIn("Is the programme full time or part time?", damaged.source_text)
        self.assertEqual(damaged.source_note, "answer appears to be a fragment of STUDENT APPLICATION #4")
        self.assertEqual(damaged.source_page, 1)

    def test_rows_only_on_the_all_sheet_wait_in_the_nearest_category(self):
        import_faq()
        only = FaqEntry.objects.filter(source_ref__startswith="all-")
        self.assertEqual(only.count(), 24)
        self.assertFalse(only.exclude(status=Status.NEEDS_CHECKING).exists())
        self.assertTrue(all(entry.source_note.startswith("Found only on the ALL sheet") for entry in only))
        defer = only.get(question="Can i defer my offer to a later intake?")
        self.assertEqual(defer.category.name, "Admission and offer letter")
        handbook = only.get(question__startswith="Where can I get an online copy of the postgraduate handbook")
        self.assertEqual(handbook.category.name, "Course registration")

    def test_repeats_on_the_all_sheet_are_not_imported(self):
        import_faq()
        rows = json.loads(default_path().read_text(encoding="utf-8"))["entries"]
        first_all = next(row for row in rows if row["sheet"] == "ALL")
        self.assertFalse(FaqEntry.objects.filter(source_ref=first_all["id"]).exists())
        self.assertEqual(FaqEntry.objects.filter(question="What is the method of application?").count(), 1)

    def test_an_unknown_category_stops_the_import(self):
        broken = {"entries": [{"sheet": "CATEGORIZED", "category": "MYSTERY", "id": "x", "question": "Q", "answer": "A"}]}
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "faq-entries.json"
            path.write_text(json.dumps(broken), encoding="utf-8")
            with self.assertRaises(ImportProblem):
                import_faq(path)
        self.assertFalse(FaqEntry.objects.exists())

    def test_the_command_reports_the_counts(self):
        out = StringIO()
        call_command("import_faq", stdout=out)
        self.assertIn("Published: 56. Needs checking: 61. Repeats on the ALL sheet skipped: 46.", out.getvalue())


class ReviewListTests(TestCase):
    def test_the_review_list_holds_every_entry_to_check(self):
        import_faq()
        content, count = review_workbook()
        self.assertEqual(count, 61)
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(sheet.max_row, 62)
        self.assertEqual(sheet["H1"].value, "Correct question (only if wrong)")
        self.assertEqual(sheet["I1"].value, "Correct answer")

    def test_text_never_runs_as_a_formula(self):
        import_faq()
        FaqEntry.objects.filter(source_ref="student-application-categorized-10").update(answer="=HYPERLINK(\"x\")")
        sheet = load_workbook(BytesIO(review_workbook()[0])).active
        answers = [row[4].value for row in sheet.iter_rows(min_row=2)]
        self.assertIn("'=HYPERLINK(\"x\")", answers)
