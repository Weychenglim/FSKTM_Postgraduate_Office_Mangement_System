"""The office's side of the chatbot (C4): student questions, numbers,
settings and loading the office's corrections."""
import tempfile
from datetime import timedelta
from io import BytesIO, StringIO
from pathlib import Path

from django.core.management import call_command
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import User

from .corrections import CorrectionsProblem, load_corrections
from .encoder import reset_encoder
from .models import FaqAltQuestion, FaqCategory, FaqChange, FaqEntry, FaqQuestion, FaqSettings
from .questions import group_key
from .review_export import review_workbook

Status = FaqEntry.Status
Outcome = FaqQuestion.Outcome


@override_settings(FAQ_ENCODER="faq.fake_encoder.FakeEncoder")
class OfficeBase(APITestCase):
    def setUp(self):
        reset_encoder()
        self.addCleanup(reset_encoder)
        self.office = User.objects.create_user(
            email="faq-desk@example.test", password="pw", full_name="Faq Desk", role=User.Role.OFFICE_ADMIN, is_staff=True
        )
        self.client.force_authenticate(self.office)
        self.category = FaqCategory.objects.create(name="During the semester", position=1)
        self.parking = FaqEntry.objects.create(question="Where can I park?", answer="At the faculty car park.", category=self.category, status=Status.PUBLISHED)

    def asked(self, text, outcome=Outcome.NO_MATCH, helpful=None, days_ago=0, typed=True, **fields):
        question = FaqQuestion.objects.create(text=text, outcome=outcome, helpful=helpful, typed=typed, **fields)
        if days_ago:
            FaqQuestion.objects.filter(pk=question.pk).update(asked_at=timezone.now() - timedelta(days=days_ago))
        return question


class StudentQuestionTests(OfficeBase):
    def test_questions_that_need_a_look_are_grouped(self):
        self.asked("Where to park?")
        self.asked("park where")
        self.asked("Parking for students?", outcome=Outcome.ANSWERED, helpful=False, entry=self.parking)
        self.asked("How do I register?", outcome=Outcome.ANSWERED)
        self.asked("hello", review=FaqQuestion.Review.DISMISSED)
        groups = self.client.get("/api/faq/questions/").data["groups"]
        self.assertEqual([(group["text"], group["count"]) for group in groups], [("park where", 2), ("Parking for students?", 1)])
        self.assertEqual((groups[0]["noMatch"], groups[0]["notHelpful"]), (2, 0))
        self.assertEqual((groups[1]["noMatch"], groups[1]["notHelpful"]), (0, 1))
        self.assertEqual((groups[0]["lastOutcome"], groups[0]["lastEntry"]), ("NO_MATCH", None))
        self.assertEqual(groups[1]["lastEntry"], {"id": str(self.parking.pk), "question": "Where can I park?"})
        self.assertEqual(group_key("Where to PARK?"), group_key("park, where"))

    def test_questions_close_in_meaning_are_grouped(self):
        self.asked("nak tangguh tawaran")
        self.asked("Can I defer my offer letter?")
        self.asked("Where to park?")
        groups = self.client.get("/api/faq/questions/").data["groups"]
        self.assertEqual([(group["text"], group["count"]) for group in groups], [("Can I defer my offer letter?", 2), ("Where to park?", 1)])
        self.assertNotIn("vector", groups[0])

    @override_settings(FAQ_SEARCH="keywords")
    def test_without_the_model_only_the_same_words_group(self):
        reset_encoder()
        self.addCleanup(reset_encoder)
        self.asked("nak tangguh tawaran")
        self.asked("Can I defer my offer letter?")
        self.assertEqual(len(self.client.get("/api/faq/questions/").data["groups"]), 2)

    def test_adding_a_group_as_another_way_to_ask(self):
        first = self.asked("tempat letak kereta?")
        second = self.asked("Tempat letak kereta")
        response = self.client.post(
            "/api/faq/questions/review/",
            {"ids": [str(first.pk), str(second.pk)], "action": "add", "entryId": str(self.parking.pk), "text": "tempat letak kereta"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["settled"], 2)
        self.assertEqual(list(FaqAltQuestion.objects.filter(entry=self.parking).values_list("text", flat=True)), ["tempat letak kereta"])
        self.assertEqual(set(FaqQuestion.objects.values_list("review", flat=True)), {"ADDED"})
        self.assertEqual(FaqQuestion.objects.get(pk=first.pk).reviewed_by, self.office)
        self.assertEqual(FaqChange.objects.get(entry=self.parking).details, {"fields": ["other ways students ask"]})
        self.assertEqual(self.client.get("/api/faq/questions/").data["groups"], [])
        again = self.client.post("/api/faq/questions/review/", {"ids": [str(first.pk)], "action": "dismiss"}, format="json")
        self.assertEqual(again.status_code, status.HTTP_409_CONFLICT)

    def test_create_and_dismiss(self):
        missing = self.asked("Is there a prayer room?")
        no_entry = self.client.post("/api/faq/questions/review/", {"ids": [str(missing.pk)], "action": "create"}, format="json")
        self.assertEqual(no_entry.status_code, status.HTTP_400_BAD_REQUEST)
        created = self.client.post(
            "/api/faq/questions/review/", {"ids": [str(missing.pk)], "action": "create", "entryId": str(self.parking.pk)}, format="json"
        )
        self.assertEqual(created.status_code, status.HTTP_200_OK)
        self.assertEqual(FaqQuestion.objects.get(pk=missing.pk).review, "CREATED")
        other = self.asked("asdf")
        self.client.post("/api/faq/questions/review/", {"ids": [str(other.pk)], "action": "dismiss"}, format="json")
        self.assertEqual(FaqQuestion.objects.get(pk=other.pk).review, "DISMISSED")

    def test_the_four_numbers(self):
        self.asked("a", outcome=Outcome.ANSWERED, helpful=True)
        self.asked("b", outcome=Outcome.ANSWERED, helpful=False)
        self.asked("c", outcome=Outcome.NO_MATCH)
        self.asked("d", outcome=Outcome.CHOICES)
        self.asked("clicked", outcome=Outcome.ANSWERED, typed=False)
        self.asked("old", outcome=Outcome.NO_MATCH, days_ago=40)
        stats = self.client.get("/api/faq/questions/").data["stats"]
        self.assertEqual(stats, {"thisMonth": 4, "answeredPercent": 50, "helpfulPercent": 50, "waiting": 3})

    def test_the_numbers_alone(self):
        self.asked("c", outcome=Outcome.NO_MATCH)
        response = self.client.get("/api/faq/questions/?only=stats")
        self.assertEqual(response.data, {"stats": {"thisMonth": 1, "answeredPercent": 0, "helpfulPercent": None, "waiting": 1}})

    def test_numbers_with_nothing_asked(self):
        self.assertEqual(
            self.client.get("/api/faq/questions/").data["stats"],
            {"thisMonth": 0, "answeredPercent": None, "helpfulPercent": None, "waiting": 0},
        )


class SettingsTests(OfficeBase):
    def test_contact_text_and_common_questions(self):
        response = self.client.patch(
            "/api/faq/settings/",
            {"contactText": "  Email pgoffice@example.test  ", "commonQuestions": ["How do I register?", "  ", "When   is add/drop?"]},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["contactText"], "Email pgoffice@example.test")
        self.assertEqual(response.data["commonQuestions"], ["How do I register?", "When is add/drop?"])
        self.assertEqual(response.data["updatedBy"], "Faq Desk")
        too_many = self.client.patch("/api/faq/settings/", {"commonQuestions": ["a", "b", "c", "d", "e"]}, format="json")
        self.assertEqual(too_many.status_code, status.HTTP_400_BAD_REQUEST)
        student = User.objects.create_user(email="faq-s@example.test", password="pw", full_name="S", role=User.Role.STUDENT)
        self.client.force_authenticate(student)
        start = self.client.get("/api/faq/start/").data
        self.assertEqual((start["contactText"], start["commonQuestions"]), ("Email pgoffice@example.test", ["How do I register?", "When is add/drop?"]))

    def test_only_the_office(self):
        for role in (User.Role.STUDENT, User.Role.COORDINATOR, User.Role.LECTURER):
            self.client.force_authenticate(User.objects.create_user(email=f"o-{role.split()[0]}@example.test", password="pw", full_name=role, role=role))
            for method, path in (("get", "/api/faq/questions/"), ("post", "/api/faq/questions/review/"), ("get", "/api/faq/settings/"), ("patch", "/api/faq/settings/")):
                with self.subTest(role=role, path=path, method=method):
                    self.assertEqual(getattr(self.client, method)(path, {"contactText": "x"}, format="json").status_code, status.HTTP_403_FORBIDDEN)
        self.assertNotEqual(FaqSettings.load().contact_text, "x")


class CorrectionsTests(OfficeBase):
    def setUp(self):
        super().setUp()
        self.wrong = FaqEntry.objects.create(question="Is it full time?", answer="Respective embassy", category=self.category, status=Status.NEEDS_CHECKING, source_ref="r1")
        self.garbled = FaqEntry.objects.create(question="rWeghiastt?", answer="x", category=self.category, status=Status.NEEDS_CHECKING, source_ref="r2")
        self.untouched = FaqEntry.objects.create(question="Weekend classes?", answer="cut", category=self.category, status=Status.NEEDS_CHECKING, source_ref="r3")

    def filled(self, fills, extra_rows=()):
        book = load_workbook(BytesIO(review_workbook()[0]))
        sheet = book.active
        for extra in extra_rows:
            sheet.append(extra)
        for row in sheet.iter_rows(min_row=2):
            number = row[0].value
            if number in fills:
                question, answer = fills[number]
                row[7].value = question
                row[8].value = answer
        folder = Path(tempfile.mkdtemp())
        path = folder / "faq-review-filled.xlsx"
        book.save(path)
        self.addCleanup(lambda: (path.unlink(), folder.rmdir()))
        return path

    def test_corrections_are_loaded_and_published(self):
        path = self.filled(
            {
                self.wrong.pk: ("", "Both. Full time and part time are offered."),
                self.garbled.pk: ("What subjects do I take?", "'=See the programme structure."),
            },
            extra_rows=[(999999, "", "", "", "", "", "", "", "Nobody")],
        )
        counts = load_corrections(path, actor=self.office)
        self.assertEqual(counts["corrected"], 2)
        self.assertEqual(counts["skipped"], 1)
        self.assertEqual(counts["problems"], ["Line 5: no entry number 999999."])
        self.wrong.refresh_from_db()
        self.assertEqual((self.wrong.status, self.wrong.answer, self.wrong.checked_by), (Status.PUBLISHED, "Both. Full time and part time are offered.", self.office))
        self.garbled.refresh_from_db()
        self.assertEqual((self.garbled.question, self.garbled.answer), ("What subjects do I take?", "=See the programme structure."))
        self.untouched.refresh_from_db()
        self.assertEqual(self.untouched.status, Status.NEEDS_CHECKING)
        actions = list(FaqChange.objects.filter(entry=self.garbled).order_by("pk").values_list("action", flat=True))
        self.assertEqual(actions, [FaqChange.Action.EDITED, FaqChange.Action.PUBLISHED])

    def test_draft_option_and_command(self):
        path = self.filled({self.wrong.pk: ("", "Both.")})
        out = StringIO()
        call_command("import_faq_corrections", str(path), "--draft", stdout=out)
        self.assertIn("Corrected: 1. Left blank: 2.", out.getvalue())
        self.wrong.refresh_from_db()
        self.assertEqual(self.wrong.status, Status.DRAFT)

    def test_a_sheet_without_the_columns_is_refused(self):
        folder = Path(tempfile.mkdtemp())
        path = folder / "wrong.xlsx"
        book = load_workbook(BytesIO(review_workbook()[0]))
        book.active["I1"].value = "Something else"
        book.save(path)
        try:
            with self.assertRaises(CorrectionsProblem):
                load_corrections(path)
        finally:
            path.unlink()
            folder.rmdir()
