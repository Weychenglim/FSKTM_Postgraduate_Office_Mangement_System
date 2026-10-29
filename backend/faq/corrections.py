"""Loading the office's corrections from the review list (export_faq_review).

Each row with a correct answer updates its entry, which then counts as
checked. Rows left blank are skipped, so the office can return the list in
parts. By default a corrected entry is published at once, because the answer
is the office's own; ``draft=True`` leaves it as a draft instead.
"""
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .entries import answer_problem, record
from .models import FaqChange, FaqEntry

ENTRY = "Entry number"
QUESTION = "Correct question (only if wrong)"
ANSWER = "Correct answer"


class CorrectionsProblem(Exception):
    pass


def _rows(path):
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise CorrectionsProblem(f"Could not open {path}: {exc}") from exc
    try:
        rows = list(book.active.iter_rows(values_only=True))
    finally:
        # A read-only workbook keeps the file open until it is closed.
        book.close()
    headings = [str(value or "").strip() for value in (rows[0] if rows else ())]
    missing = [name for name in (ENTRY, QUESTION, ANSWER) if name not in headings]
    if missing:
        raise CorrectionsProblem(f"The review list is missing these columns: {', '.join(missing)}")
    at = {name: headings.index(name) for name in (ENTRY, QUESTION, ANSWER)}
    for line, row in enumerate(rows[1:], start=2):
        cells = list(row) + [None] * len(headings)
        yield line, cells[at[ENTRY]], str(cells[at[QUESTION]] or "").strip(), str(cells[at[ANSWER]] or "").strip()


def _text(value):
    # Text that looked like a formula was written with a leading apostrophe.
    return value[1:] if value.startswith("'") else value


@transaction.atomic
def load_corrections(path, *, actor=None, draft=False):
    counts = {"corrected": 0, "skipped": 0, "problems": []}
    for line, number, question, answer in _rows(path):
        if not answer:
            counts["skipped"] += 1
            continue
        entry = FaqEntry.objects.select_for_update().filter(pk=number).first() if str(number or "").isdigit() else None
        if entry is None:
            counts["problems"].append(f"Line {line}: no entry number {number}.")
            continue
        answer = _text(answer)
        problem = answer_problem(answer)
        if problem:
            counts["problems"].append(f"Line {line}: {problem}")
            continue
        changed = ["answer"] if answer != entry.answer else []
        entry.answer = answer
        if question and _text(question) != entry.question:
            entry.question = " ".join(_text(question).split())[:500]
            changed.append("question")
        if not entry.question.strip():
            counts["problems"].append(f"Line {line}: entry {number} has no question. Fill in the correct question.")
            continue
        previous = entry.status
        if previous == FaqEntry.Status.NEEDS_CHECKING:
            entry.checked_by = actor
            entry.checked_at = timezone.now()
        entry.status = FaqEntry.Status.DRAFT if draft else FaqEntry.Status.PUBLISHED
        entry.updated_by = actor
        entry.save()
        record(entry, actor, FaqChange.Action.EDITED, {"fields": changed, "source": "the office's review list"})
        if previous != entry.status:
            action = FaqChange.Action.PUBLISHED if entry.status == FaqEntry.Status.PUBLISHED else FaqChange.Action.CHECKED
            record(entry, actor, action, {"from": previous})
        counts["corrected"] += 1
    return counts
