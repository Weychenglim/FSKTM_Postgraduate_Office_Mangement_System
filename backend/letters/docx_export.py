"""The Word (.docx) copy of an issued letter, laid out like the printed one.

The letterhead text mirrors ``LETTERHEAD`` in frontend/src/utils/letterDocument.ts;
a frontend test fails if the two drift apart. The body uses the same light
markup as the template editor: ``**bold**``, ``*italic*``, ``__underline__``,
blank lines between paragraphs and ``- `` bullet lines.
"""
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

LOGO = Path(__file__).resolve().parent / "assets" / "um-logo.png"

FACULTY = "Faculty of Computer Science & Information Technology"
ADDRESS_LINES = (
    "Academic Affairs Office, Level 4",
    "50603 Kuala Lumpur, Malaysia",
    "+603-7967 6300  |  pg.fsktm@um.edu.my",
)
DISCLAIMER = (
    "This letter is issued upon the student's request for official purposes only. "
    "Any alteration to this document shall render it invalid."
)
SIGNATORY_NAME = "Dean of Postgraduate Studies"
SIGNATORY_OFFICE = "FSKTM Office Executive Secretariat"

TEXT_WIDTH = Mm(178)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0x4F, 0x46, 0xE5)

INLINE_RE = re.compile(r"\*\*([^*]+?)\*\*|__([^_]+?)__|\*([^*\n]+?)\*")
BULLET_RE = re.compile(r"^\s*-\s+")


def _runs(text, bold=False, italic=False, underline=False):
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            yield text[position:match.start()], bold, italic, underline
        strong, underlined, emphasis = match.groups()
        if strong is not None:
            yield from _runs(strong, True, italic, underline)
        elif underlined is not None:
            yield from _runs(underlined, bold, italic, True)
        else:
            yield from _runs(emphasis, bold, True, underline)
        position = match.end()
    if position < len(text):
        yield text[position:], bold, italic, underline


def _add_marked_text(paragraph, text):
    for chunk, bold, italic, underline in _runs(text):
        run = paragraph.add_run(chunk)
        run.bold = bold or None
        run.italic = italic or None
        run.underline = underline or None


def _styled_run(paragraph, text, *, size, color=None, bold=False):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold or None
    if color is not None:
        run.font.color.rgb = color
    return run


def _bottom_rule(paragraph):
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "0C1424")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _letterhead(document):
    table = document.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    brand = left.paragraphs[0]
    if LOGO.exists():
        brand.add_run().add_picture(str(LOGO), height=Mm(12))
    faculty = left.add_paragraph()
    _styled_run(faculty, FACULTY.upper(), size=7, color=MUTED, bold=True)
    address = right.paragraphs[0]
    address.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for index, line in enumerate(ADDRESS_LINES):
        run = _styled_run(address, line, size=7.5, color=MUTED)
        if index < len(ADDRESS_LINES) - 1:
            run.add_break()
    rule = document.add_paragraph()
    _bottom_rule(rule)
    rule.paragraph_format.space_after = Pt(14)


def _meta_line(document, letter):
    meta = document.add_paragraph()
    meta.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    meta.paragraph_format.space_after = Pt(16)
    _styled_run(meta, "Our Ref: ", size=9, color=MUTED)
    _styled_run(meta, letter.reference_number, size=9, bold=True)
    _styled_run(meta, "\tDate: ", size=9, color=MUTED)
    _styled_run(meta, letter.issued_on.strftime("%d %B %Y"), size=9, bold=True)


def _body(document, content):
    for block in re.split(r"\n{2,}", content.replace("\r\n", "\n").strip()):
        lines = block.split("\n")
        filled = [line for line in lines if line.strip()]
        if not filled:
            continue
        if all(BULLET_RE.match(line) for line in filled):
            for line in filled:
                _add_marked_text(
                    document.add_paragraph(style="List Bullet"), BULLET_RE.sub("", line)
                )
            continue
        paragraph = document.add_paragraph()
        for index, line in enumerate(lines):
            if index:
                paragraph.add_run().add_break()
            _add_marked_text(paragraph, line)


def _closing(document):
    disclaimer = document.add_paragraph()
    _styled_run(disclaimer, DISCLAIMER, size=8.5, color=MUTED)
    signature = document.add_paragraph()
    signature.paragraph_format.space_before = Pt(42)
    signature.paragraph_format.keep_with_next = True
    _styled_run(signature, "_" * 28, size=10, color=MUTED)
    name = document.add_paragraph()
    name.paragraph_format.space_after = Pt(0)
    name.paragraph_format.keep_with_next = True
    _styled_run(name, SIGNATORY_NAME, size=10.5, bold=True)
    office = document.add_paragraph()
    office.paragraph_format.keep_with_next = True
    _styled_run(office, SIGNATORY_OFFICE.upper(), size=7, color=MUTED, bold=True)
    stamp = document.add_paragraph()
    _styled_run(stamp, "Computer-generated. Verify with the PG Office.", size=7.5, color=ACCENT, bold=True)


def build_letter_docx(letter):
    """Return the .docx bytes for ``letter``."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(16)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK

    properties = document.core_properties
    properties.title = f"{letter.template_name} {letter.reference_number}"
    properties.author = "FSKTM Postgraduate Office"

    _letterhead(document)
    _meta_line(document, letter)
    _body(document, letter.content)
    _closing(document)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
