from io import BytesIO
from openpyxl import load_workbook
from django.test import override_settings
from rest_framework.test import APITestCase

from . import test_appointment_lifecycle as fixture
from .models import SupervisorApplication, ResearchAmendment
from dashboard.reports import build_workflow_report, build_workflow_report_workbook
from dashboard.dossiers import build_student_progress_dossier


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class StudentCancellationVisibilityTests(APITestCase):
    def setUp(self):
        fixture.AppointmentLifecycleTests.setUp(self)
        self.office.is_staff = True
        self.office.save(update_fields=['is_staff'])
    _lecturer = fixture.AppointmentLifecycleTests._lecturer

    def make_application(self, status):
        return SupervisorApplication.objects.create(
            student=self.student, proposed_supervisor=self.supervisor,
            academic_semester=self.semester, status=status,
            research_title='Withdrawn proposal', research_abstract='Proposal abstract',
        )

    def research(self, cancel_actor):
        self.client.force_authenticate(self.student_user)
        response = self.client.post('/api/appointments/research-amendments/', {
            'kind': 'RESEARCH', 'title': 'Revised title', 'reason': 'Research update',
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        pk = response.data['id']
        if cancel_actor is None:
            return pk
        self.client.force_authenticate(cancel_actor)
        response = self.client.post(f'/api/appointments/research-amendments/{pk}/cancel/', {'reason': 'Withdraw request'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        return pk

    def test_supervisor_student_cancellation_hidden_from_staff_lists_and_reports(self):
        withdrawn = self.make_application('CANCELLED_BY_STUDENT')
        office_cancelled = self.make_application('CANCELLED_BY_OFFICE')
        for actor, path, id_key in [
            (self.supervisor, 'request-history', 'applicationId'),
            (self.coordinator, 'coordinator-history', 'id'),
            (self.office, '', 'applicationId'),
        ]:
            self.client.force_authenticate(actor)
            response = self.client.get('/api/appointments/supervisor/' + (f'{path}/' if path else ''))
            self.assertEqual(response.status_code, 200, response.data)
            ids = {row[id_key] for row in response.data}
            self.assertNotIn(withdrawn.pk, ids)
            self.assertIn(office_cancelled.pk, ids)
            report = build_workflow_report(actor, {'semester': 'all'})
            ids = {str(row['recordId']) for row in report['supervisor']['records']}
            self.assertNotIn(str(withdrawn.pk), ids)
            self.assertIn(str(office_cancelled.pk), ids)
            dossier = build_student_progress_dossier(actor, self.student.matric_no)
            self.assertNotIn(str(withdrawn.pk), {str(row['recordId']) for row in dossier['supervisor']['records']})
            self.assertIn(str(office_cancelled.pk), {str(row['recordId']) for row in dossier['supervisor']['records']})
            self.assertEqual(report['supervisor']['total'], len(report['supervisor']['records']))
            sheet = load_workbook(BytesIO(build_workflow_report_workbook(report)))['Supervisor']
            exported_ids = {str(row[0]) for row in list(sheet.values)[1:]}
            self.assertNotIn(str(withdrawn.pk), exported_ids)
            self.assertIn(str(office_cancelled.pk), exported_ids)
            team = self.client.get(f'/api/appointments/co-supervisor/students/{self.student.pk}/').data
            self.assertNotIn(f'primary-workflow-{withdrawn.pk}', {event['id'] for event in team['timeline']})
        self.client.force_authenticate(self.student_user)
        self.assertIn(withdrawn.pk, {row['id'] for row in self.client.get('/api/appointments/supervisor/applications/').data})
        team = self.client.get(f'/api/appointments/co-supervisor/students/{self.student.pk}/').data
        self.assertIn(f'primary-workflow-{withdrawn.pk}', {event['id'] for event in team['timeline']})
        self.client.force_authenticate(self.office)
        self.assertEqual(self.client.get(f'/api/appointments/supervisor/applications/{withdrawn.pk}/').status_code, 200, 'Authorised audit detail remains available')
        self.assertTrue(SupervisorApplication.objects.filter(pk=self.application.pk).exists())

    def test_research_cancellation_uses_cancellation_actor_not_request_owner(self):
        withdrawn = self.research(self.student_user)
        office_cancelled = self.research(self.office)
        system_cancelled = self.research(None)
        from .research_amendments import cancel_pending_amendments
        cancel_pending_amendments(self.student.pk, self.office, 'Automatic lifecycle cleanup')
        for actor in (self.supervisor, self.coordinator, self.office):
            self.client.force_authenticate(actor)
            response = self.client.get('/api/appointments/research-amendments/')
            ids = {row['id'] for row in response.data['requests']}
            self.assertNotIn(withdrawn, ids)
            self.assertIn(office_cancelled, ids, 'Office cancellation of a Student-owned request remains visible')
            self.assertIn(system_cancelled, ids)
            report = build_workflow_report(actor, {'semester': 'all'})
            ids = {str(row['recordId']) for row in report['researchAmendments']['records']}
            self.assertNotIn(str(withdrawn), ids)
            self.assertIn(str(office_cancelled), ids)
            self.assertIn(str(system_cancelled), ids)
            dossier = build_student_progress_dossier(actor, self.student.matric_no)
            self.assertNotIn(withdrawn, {row['id'] for row in dossier['researchAmendments']})
            sheet = load_workbook(BytesIO(build_workflow_report_workbook(report)))['Research Amendments']
            exported_ids = {str(row[0]) for row in list(sheet.values)[1:]}
            self.assertNotIn(str(withdrawn), exported_ids)
            self.assertIn(str(office_cancelled), exported_ids)
            self.assertEqual(self.client.get(f'/api/appointments/research-amendments/{withdrawn}/').status_code, 200)
        self.client.force_authenticate(self.student_user)
        self.assertIn(withdrawn, {row['id'] for row in self.client.get('/api/appointments/research-amendments/').data['requests']})
        self.assertTrue(ResearchAmendment.objects.get(pk=withdrawn).events.filter(action='CANCEL').exists())
        self.profile.refresh_from_db()
        self.assertNotEqual(self.profile.proposed_topic, 'Revised title')
        self.client.force_authenticate(self.new_supervisor)
        self.assertEqual(self.client.get(f'/api/appointments/research-amendments/{withdrawn}/').status_code, 404, 'Audit access does not widen scope')
