from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

from django.db import connection
from django.test import override_settings
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from academics.models import AcademicSemester
from accounts.models import CoordinatorDelegation
from . import test_coordinator_workspace as fixture
from .models import CoSupervisorAppointment, SupervisorAppointment


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class SupervisionExportTests(APITestCase):
    _lecturer = fixture.CoordinatorWorkspaceTests._lecturer
    application_for = fixture.CoordinatorWorkspaceTests.application_for
    nomination = fixture.CoordinatorWorkspaceTests.nomination

    def setUp(self):
        fixture.CoordinatorWorkspaceTests.setUp(self)
        self.office.is_staff = True
        self.office.save(update_fields=["is_staff"])

    def get(self, actor=None, namespace="office-workspace", **params):
        self.client.force_authenticate(actor or self.office)
        return self.client.get(f"/api/appointments/{namespace}/supervision-export/", params)

    def workbook(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment;", response["Content-Disposition"])
        return load_workbook(BytesIO(response.content))

    def records(self, workbook, sheet):
        rows = list(workbook[sheet].values)
        return [dict(zip(rows[0], row)) for row in rows[1:]]

    def appoint(self, number, **kwargs):
        application = self.application_for(number, status="APPROVED", **kwargs)
        return SupervisorAppointment.objects.create(application=application, student=application.student,
            supervisor=self.supervisor, approved_by=self.coordinator)

    def support(self, supervisor=None, status="ACTIVE"):
        supervisor = supervisor or self.new_supervisor
        return CoSupervisorAppointment.objects.create(nomination=self.nomination(supervisor, "APPROVED"),
            student=self.student, supervisor=supervisor, approved_by=self.coordinator, status=status)

    def test_only_active_appointments_export_actual_roles_and_complete_team(self):
        self.support()
        self.support(self.new_panel)
        self.support(self.panel, status="ENDED")
        self.nomination(self.other_coordinator)
        ended = self.appoint(1)
        ended.status = "ENDED"
        ended.save(update_fields=["status"])
        self.application_for(2)
        workbook = self.workbook(self.get())
        self.assertEqual(workbook.sheetnames, ["Students by lecturer", "Supervisors by student"])
        relationships = self.records(workbook, "Students by lecturer")
        self.assertEqual({(row["Lecturer name"], row["Role"]) for row in relationships},
                         {("OLD-001", "Primary supervisor"), ("NEW-001", "Supporting supervisor"),
                          ("PANEL-002", "Supporting supervisor")})
        self.assertEqual({row["Matric number"] for row in relationships}, {"DEMO-LIFE-001"})
        students = self.records(workbook, "Supervisors by student")
        self.assertEqual(len(students), 1)
        self.assertEqual(students[0]["Active primary supervisor"], "OLD-001")
        self.assertEqual(students[0]["All active supporting supervisors"], "NEW-001; PANEL-002")

    def test_selected_lecturer_limits_relationships_but_keeps_students_complete_current_team(self):
        self.support()
        self.support(self.new_panel)
        workbook = self.workbook(self.get(supervisor=self.new_supervisor.pk))
        relationships = self.records(workbook, "Students by lecturer")
        self.assertEqual(len(relationships), 1)
        self.assertEqual(relationships[0]["Staff number"], "NEW-001")
        student = self.records(workbook, "Supervisors by student")[0]
        self.assertEqual(student["Active primary supervisor"], "OLD-001")
        self.assertEqual(student["All active supporting supervisors"], "NEW-001; PANEL-002")

    def test_export_uses_all_matched_pages_and_queries_do_not_grow_per_student(self):
        with CaptureQueriesContext(connection) as small:
            self.workbook(self.get())
        for number in range(1, 28):
            self.appoint(number)
        with CaptureQueriesContext(connection) as large:
            workbook = self.workbook(self.get(page=2, size=10))
        self.assertEqual(len(self.records(workbook, "Students by lecturer")), 28)
        self.assertEqual(len(self.records(workbook, "Supervisors by student")), 28)
        self.assertLessEqual(len(large) - len(small), 1)
        self.assertLessEqual(len(large), 10)

    def test_namespace_guards_reject_other_roles_inactive_and_nonstaff_office(self):
        for namespace, actors in (
            ("office-workspace", (self.coordinator, self.supervisor, self.student_user)),
            ("coordinator-workspace", (self.office, self.supervisor, self.student_user)),
        ):
            for actor in actors:
                with self.subTest(namespace=namespace, role=actor.role):
                    self.assertEqual(self.get(actor, namespace).status_code, 403)
        self.office.is_staff = False
        self.office.save(update_fields=["is_staff"])
        self.assertEqual(self.get().status_code, 403)
        self.office.is_staff, self.office.is_active = True, False
        self.office.save(update_fields=["is_staff", "is_active"])
        self.assertEqual(self.get().status_code, 403)
        self.coordinator.is_active = False
        self.coordinator.save(update_fields=["is_active"])
        self.assertEqual(self.get(self.coordinator, "coordinator-workspace").status_code, 403)
        self.client.force_authenticate(None)
        self.assertIn(self.client.get("/api/appointments/office-workspace/supervision-export/").status_code, (401, 403))

    def test_coordinator_scope_includes_only_managed_and_current_delegated_programmes(self):
        self.appoint(1, programme="Acting Programme")
        self.appoint(2, programme="Outside")
        today = timezone.localdate()
        CoordinatorDelegation.objects.create(coordinator=self.coordinator, programme="Acting Programme",
            starts_on=today, ends_on=today, justification="Leave", granted_by=self.office)
        workbook = self.workbook(self.get(self.coordinator, "coordinator-workspace"))
        self.assertEqual({row["Matric number"] for row in self.records(workbook, "Supervisors by student")},
                         {"DEMO-LIFE-001", "BROWSE-001"})
        with patch("accounts.delegations.today", return_value=today + timedelta(days=1)):
            workbook = self.workbook(self.get(self.coordinator, "coordinator-workspace"))
        self.assertEqual({row["Matric number"] for row in self.records(workbook, "Supervisors by student")},
                         {"DEMO-LIFE-001"})
        office = self.workbook(self.get())
        self.assertEqual(len(self.records(office, "Supervisors by student")), 3)

    def test_team_relationship_filters_search_and_student_membership_match_browse(self):
        today = timezone.localdate()
        second = AcademicSemester.objects.create(academic_session=self.semester.academic_session,
            term=AcademicSemester.Term.SEMESTER_II, starts_on=today + timedelta(days=150),
            ends_on=today + timedelta(days=240), created_by=self.office)
        self.supervisor.full_name = "DrOne"
        self.supervisor.save(update_fields=["full_name"])
        self.new_supervisor.full_name = "DrTwo"
        self.new_supervisor.save(update_fields=["full_name"])
        self.supervisor_appointment.appointment_date = today - timedelta(days=30)
        self.supervisor_appointment.save(update_fields=["appointment_date"])
        supporting = self.support()
        supporting.appointment_date = today - timedelta(days=2)
        supporting.save(update_fields=["appointment_date"])
        supporting.nomination.academic_semester = second
        supporting.nomination.save(update_fields=["academic_semester"])
        cases = (
            ({"semester": self.semester.label}, {"DrOne"}),
            ({"semester": second.label}, {"DrTwo"}),
            ({"type": "primary"}, {"DrOne"}),
            ({"type": "supporting"}, {"DrTwo"}),
            ({"from": (today - timedelta(days=5)).isoformat()}, {"DrTwo"}),
            ({"to": (today - timedelta(days=10)).isoformat()}, {"DrOne"}),
            ({"q": "DrOne"}, {"DrOne"}),
            ({"q": "Lifecycle-safe"}, {"DrOne", "DrTwo"}),
            ({"q": self.student.matric_no}, {"DrOne", "DrTwo"}),
            ({"q": self.student.user.full_name}, {"DrOne", "DrTwo"}),
            ({"q": "DrOne", "view": "students"}, {"DrOne", "DrTwo"}),
            ({"q": "DrOne", "supervisor": self.new_supervisor.pk}, set()),
            ({"programme": "  master of artificial intelligence (coursework) "}, {"DrOne", "DrTwo"}),
            ({"programme": "Outside"}, set()),
            ({"status": "ENDED"}, set()),
            ({"q": "not found"}, set()),
        )
        for params, names in cases:
            with self.subTest(params=params):
                workbook = self.workbook(self.get(**params))
                self.assertEqual({row["Lecturer name"] for row in self.records(workbook, "Students by lecturer")}, names)
                browse = self.client.get("/api/appointments/office-workspace/teams/", {**params, "view": "students"})
                self.assertEqual({row["Matric number"] for row in self.records(workbook, "Supervisors by student")},
                                 {row["matricNo"] for row in browse.data["results"]})
                if names:
                    student = self.records(workbook, "Supervisors by student")[0]
                    self.assertEqual(student["Active primary supervisor"], "DrOne")
                    self.assertEqual(student["All active supporting supervisors"], "DrTwo")

    def test_empty_and_supporting_only_student_do_not_substitute_pending_primary(self):
        self.supervisor_appointment.status = "ENDED"
        self.supervisor_appointment.save(update_fields=["status"])
        self.nomination()
        self.assertEqual(len(self.records(self.workbook(self.get()), "Supervisors by student")), 0)
        self.support()
        row = self.records(self.workbook(self.get()), "Supervisors by student")[0]
        self.assertIsNone(row["Active primary supervisor"])
        self.assertEqual(row["All active supporting supervisors"], "NEW-001")

    def test_supervisor_name_search_must_match_a_relationship_that_passes_team_filters(self):
        today = timezone.localdate()
        second = AcademicSemester.objects.create(academic_session=self.semester.academic_session,
            term=AcademicSemester.Term.SEMESTER_II, starts_on=today + timedelta(days=150),
            ends_on=today + timedelta(days=240), created_by=self.office)
        self.supervisor_appointment.appointment_date = today - timedelta(days=30)
        self.supervisor_appointment.save(update_fields=["appointment_date"])
        supporting = self.support()
        supporting.nomination.academic_semester = second
        supporting.nomination.save(update_fields=["academic_semester"])
        cases = ({"q": "OLD-001", "type": "supporting"},
                 {"q": "OLD-001", "semester": second.label},
                 {"q": "OLD-001", "from": (today - timedelta(days=5)).isoformat()},
                 {"q": "OLD-001", "supervisor": self.new_supervisor.pk})
        for params in cases:
            for view in ("supervisors", "students"):
                with self.subTest(params=params, view=view):
                    workbook = self.workbook(self.get(**params, view=view))
                    self.assertEqual(self.records(workbook, "Students by lecturer"), [])
                    self.assertEqual(self.records(workbook, "Supervisors by student"), [])
                    browse = self.client.get("/api/appointments/office-workspace/teams/", {**params, "view": view})
                    self.assertEqual(browse.data["total"], 0)
        for query in (self.student.matric_no, self.student.user.full_name, "Lifecycle-safe", "NEW-001"):
            with self.subTest(query=query):
                workbook = self.workbook(self.get(q=query, type="supporting", semester=second.label))
                relationships = self.records(workbook, "Students by lecturer")
                self.assertEqual([row["Lecturer name"] for row in relationships], ["NEW-001"])
                students = self.records(workbook, "Supervisors by student")
                self.assertEqual([row["Active primary supervisor"] for row in students], ["OLD-001"])
                self.assertEqual(students[0]["All active supporting supervisors"], "NEW-001")

    def test_user_controlled_text_remains_literal_and_not_spreadsheet_formulas(self):
        self.supervisor.full_name = '=HYPERLINK("https://example.test", "click")'
        self.supervisor.save(update_fields=["full_name"])
        self.supervisor.lecturer.staff_no = "+DANGEROUS"
        self.supervisor.lecturer.save(update_fields=["staff_no"])
        self.student.user.full_name = "@DANGEROUS"
        self.student.user.save(update_fields=["full_name"])
        self.student.matric_no = "-DANGEROUS"
        self.student.programme = "=1+1"
        self.student.save(update_fields=["matric_no", "programme"])
        self.support()
        workbook = self.workbook(self.get())
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f", cell.coordinate)
                    if isinstance(cell.value, str):
                        self.assertEqual(cell.data_type, "s", cell.coordinate)
        lecturer = self.records(workbook, "Students by lecturer")[0]
        self.assertEqual(lecturer["Lecturer name"], self.supervisor.full_name)
        self.assertEqual(lecturer["Programme"], "=1+1")

    def test_invalid_team_filters_return_400(self):
        for params in ({"from": "bad"}, {"to": "2026-99-99"}, {"type": "panel"},
                       {"supervisor": "bad"}, {"view": "history"}, {"sort": "random"},
                       {"from": "2026-10-14", "to": "2026-10-13"}):
            with self.subTest(params=params):
                self.assertEqual(self.get(**params).status_code, 400)
