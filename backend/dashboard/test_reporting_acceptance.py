"""A05 persisted report/export/dossier/reconciliation combinations."""
from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.db import DatabaseError
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from accounts.delegations import grant_delegation, revoke_delegation
from accounts.participant_lifecycle import transition_student
from academics.models import AcademicSemester
from academics.services import activate_semester
from academics.test_capacity_helpers import publish_test_capacity_plan
from appointments import co_supervision, research_amendments
from appointments.models import CoSupervisorAppointment, SupervisorAppointment
from marks.closure_preview import semester_closure_preview
from marks.models import EvaluationTask, MarkEntry, MarkScore
from .acceptance_fixtures import ReportingAcceptanceFixture
from .models import WorkflowReconciliationAudit


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class ReportingAcceptanceTests(ReportingAcceptanceFixture, APITestCase):
    def get(self, actor, path, params=None, expected=200):
        self.client.force_authenticate(actor)
        response = self.client.get(path, params or {})
        self.assertEqual(response.status_code, expected, getattr(response, 'data', None))
        return response

    def report(self, actor, **filters):
        return self.get(actor, '/api/dashboard/reports/', filters or {'semester': 'all'}).data

    def workbook(self, actor, **filters):
        response = self.get(actor, '/api/dashboard/reports/export/', filters or {'semester': 'all'})
        return load_workbook(BytesIO(response.content))

    def preserved(self):
        return {model: list(model.objects.order_by('pk').values()) for model in (
            EvaluationTask, MarkEntry, MarkScore, SupervisorAppointment, CoSupervisorAppointment,
        )}

    def assert_preserved(self, before):
        self.assertEqual(self.preserved(), before)

    def issue(self):
        listing = self.get(self.office, '/api/dashboard/reconciliation/', {'pageSize': 100})
        return next(row for row in listing.data['results'] if row['recordType'] == 'SUPERVISOR_APPLICATION'
                    and row['recordId'] == str(self.legacy.pk) and row['issueType'] == 'SEMESTER_UNASSIGNED')

    def repair(self, issue, reason='A05 verified legacy semester'):
        self.client.force_authenticate(self.office)
        return self.client.post(f"/api/dashboard/reconciliation/issues/{issue['issueId']}/apply/", {
            'expectedFingerprint': issue['fingerprint'], 'reason': reason,
            'resolution': {'action': 'ASSIGN_SEMESTER', 'semesterId': self.semester.pk},
        }, format='json')

    def test_report_and_real_workbook_totals_ids_and_formula_cells_agree(self):
        before = self.preserved()
        for actor in (self.office, self.coordinator, self.supervisor, self.new_supervisor):
            with self.subTest(role=actor.role, actor=actor.pk):
                report = self.report(actor)
                workbook = self.workbook(actor)
                self.assertEqual(report['overview']['totalRecords'], sum(
                    report[section]['total'] for section in ('supervisor', 'panel', 'researchAmendments', 'marks', 'timeline')
                    if report[section] is not None))
                for section, sheet in (('supervisor', 'Supervisor'), ('panel', 'Panel'), ('marks', 'Marks'), ('timeline', 'Timeline')):
                    if report[section] is None:
                        self.assertNotIn(sheet, workbook.sheetnames)
                        continue
                    expected = [str(row['recordId']) for row in report[section]['records']]
                    self.assertEqual([str(row[0]) for row in workbook[sheet].iter_rows(min_row=2, values_only=True)], expected)
                for sheet in workbook:
                    self.assertFalse(any(cell.data_type == 'f' for row in sheet for cell in row))
                if actor == self.coordinator:
                    self.assertIsNone(report['marks'])
                    self.assertNotIn(self.foreign_student.matric_no, str(report))
        # Formula-leading names are retained as literal text in an actual exported workbook.
        self.assertTrue(any(cell.value == '=A05 foreign student' and cell.data_type == 's'
                            for row in self.workbook(self.office)['Supervisor'] for cell in row))
        self.assert_preserved(before)

    def test_student_dossier_redacts_lifecycle_marks_internal_calendar_and_other_students(self):
        transition_student(matric_no=self.student.matric_no, target_status='DEFERRED', reason='PRIVATE lifecycle decision', actor=self.office)
        before = self.preserved()
        student = self.get(self.student_user, self.dossier_url()).data
        office = self.get(self.office, self.dossier_url()).data
        self.assertEqual(student['visibility'], 'PUBLIC')
        self.assertIsNone(student['marks'])
        self.assertNotIn('MARKS', student['visibleSections'])
        self.assertNotIn('PRIVATE lifecycle decision', str(student))
        self.assertNotIn('PRIVATE Office calendar', str(student))
        self.assertIn('PRIVATE lifecycle decision', str(office))
        self.assertEqual(student['student']['status'], 'Deferred')
        self.get(self.student_user, self.dossier_url(self.foreign_student), expected=404)
        self.get(self.student_user, '/api/dashboard/reports/', expected=403)
        self.get(self.student_user, '/api/dashboard/reports/export/', expected=403)
        self.assert_preserved(before)

    def test_supporting_replacement_and_closure_keep_scoped_history_without_current_research(self):
        replacement = co_supervision.nominate(actor=self.supervisor, student_id=self.student.pk,
            candidate_id=self.new_panel.pk, justification='Replacement acceptance', replaces_appointment_id=self.supporting.pk)
        co_supervision.decide(nomination_id=replacement.pk, actor=self.new_panel, action='accept')
        co_supervision.decide(nomination_id=replacement.pk, actor=self.coordinator, action='approve')
        rows = self.report(self.office)['supervisor']['records']
        old = next(row for row in rows if row['recordType'] == 'CO_SUPERVISOR_NOMINATION' and row['recordId'] == str(self.nomination.pk))
        self.assertEqual(old['replacementAppointmentId'], replacement.appointment.pk)
        self.assertNotEqual(old['appointmentStatus'], 'ACTIVE')
        former = self.report(self.new_supervisor)
        old = next(row for row in former['supervisor']['records'] if row['recordType'] == 'CO_SUPERVISOR_NOMINATION')
        self.assertEqual(old['researchTitle'], '')
        self.assertEqual(former['marks']['total'], 0)
        before = list(MarkEntry.objects.values())
        co_supervision.end_appointment(appointment_id=replacement.appointment.pk, actor=self.office,
                                      reason='PRIVATE supporting closure', outcome='OTHER')
        current = self.report(self.office)['supervisor']['records']
        self.assertEqual(next(row for row in current if row['recordType'] == 'CO_SUPERVISOR_NOMINATION'
                              and row['recordId'] == str(replacement.pk))['appointmentEndReason'], 'PRIVATE supporting closure')
        public = self.get(self.student_user, self.dossier_url()).data
        self.assertNotIn('PRIVATE supporting closure', str(public))
        self.assertEqual(list(MarkEntry.objects.values()), before)

    def test_transfer_changes_programme_scope_atomically_and_retains_team_marks_and_own_decisions(self):
        before = self.preserved()
        target = self.foreign_student.programme
        row = research_amendments.submit(actor=self.office, data={
            'kind': 'TRANSFER', 'studentId': self.student.pk, 'destinationProgramme': target, 'reason': 'A05 programme transfer',
        })
        research_amendments.decide(request_id=row.pk, actor=self.coordinator, decision='APPROVE', retain_team=True, expected_status='PENDING_SOURCE_COORDINATOR')
        research_amendments.decide(request_id=row.pk, actor=self.other_coordinator, decision='APPROVE', retain_team=True, expected_status='PENDING_DESTINATION_COORDINATOR')
        self.student.refresh_from_db(); self.profile.refresh_from_db()
        self.assertEqual(self.student.programme, target)
        self.assertEqual(self.profile.programme, target)
        self.get(self.coordinator, self.dossier_url(), expected=404)
        self.get(self.other_coordinator, self.dossier_url())
        for section in ('supervisor', 'panel'):
            self.assertEqual(self.report(self.coordinator)[section]['total'], 0)
        historical = self.report(self.coordinator)['researchAmendments']['records']
        self.assertIn(str(row.pk), [item['recordId'] for item in historical])
        self.assertNotIn('Marks', self.workbook(self.coordinator).sheetnames)
        self.assert_preserved(before)

    def test_delegation_revocation_removes_current_report_export_and_dossier_scope(self):
        before = self.preserved()
        today = timezone.localdate()
        grant = grant_delegation(actor=self.office, coordinator_id=self.other_coordinator.pk,
            programme=self.student.programme, starts_on=today, ends_on=today + timedelta(days=1), justification='A05 temporary cover')
        active = self.report(self.other_coordinator)
        self.assertIn(self.student.matric_no, str(active['supervisor']['records']))
        self.get(self.other_coordinator, self.dossier_url())
        self.assertIsNone(active['marks'])
        revoke_delegation(delegation_id=grant.pk, actor=self.office, reason='A05 revoke cover')
        report = self.report(self.other_coordinator)
        self.assertNotIn(self.student.matric_no, str(report['supervisor']['records']))
        book = self.workbook(self.other_coordinator)
        self.assertNotIn(self.student.matric_no, str(list(book['Supervisor'].values)))
        self.get(self.other_coordinator, self.dossier_url(), expected=404)
        self.assert_preserved(before)

    def test_handover_preserves_source_semester_report_and_dossier_history(self):
        today = timezone.localdate()
        self.semester.ends_on = today - timedelta(days=1)
        self.semester.save(update_fields=['ends_on'])
        target = AcademicSemester.objects.create(code='A05-NEXT', academic_session='2100/2101', term='SEMESTER_II',
            starts_on=today, ends_on=today + timedelta(days=100), created_by=self.office)
        publish_test_capacity_plan(target, self.office)
        before = self.preserved()
        preview = semester_closure_preview(self.semester, target_semester=target)
        activate_semester(target, actor=self.office, reason='A05 handover', preview_token=preview['previewToken'], acknowledge_unfinished=True)
        active = self.report(self.office, semester='active')
        history = self.report(self.office, semester=self.semester.code)
        self.assertEqual(active['marks']['total'], 0)
        self.assertEqual(history['marks']['records'][0]['semesterCode'], self.semester.code)
        self.assertEqual(history['marks']['total'], 1)
        dossier = self.get(self.office, self.dossier_url()).data
        self.assertEqual(dossier['marks']['tasks'][0]['semesterCode'], self.semester.code)
        self.assertEqual(dossier['timeline']['semesterCode'], target.code)
        self.assert_preserved(before)

    def test_reconciliation_reason_stale_preview_replay_and_audit_scope(self):
        issue = self.issue()
        before = self.preserved()
        self.assertEqual(self.repair(issue, reason=' ').status_code, 400)
        self.assertEqual(self.repair({**issue, 'fingerprint': '0' * 64}).status_code, 409)
        self.assertEqual(WorkflowReconciliationAudit.objects.count(), 0)
        self.legacy.refresh_from_db(); self.assertIsNone(self.legacy.academic_semester_id)
        self.assertEqual(self.repair(issue).status_code, 200)
        self.assertEqual(self.repair(issue).status_code, 409)
        self.assertEqual(WorkflowReconciliationAudit.objects.count(), 1)
        audit = WorkflowReconciliationAudit.objects.get()
        self.assertEqual(audit.actor, self.office)
        self.assertEqual(audit.reason, 'A05 verified legacy semester')
        self.assertIsNone(audit.before_values['academicSemesterId'])
        self.assertEqual(audit.after_values['academicSemesterId'], self.semester.pk)
        scoped = self.get(self.office, '/api/dashboard/reconciliation/audits/', {'entityType': 'SUPERVISOR_APPLICATION', 'entityId': str(self.legacy.pk)})
        self.assertEqual([row['id'] for row in scoped.data['results']], [audit.pk])
        for actor in (self.student_user, self.supervisor, self.coordinator):
            self.get(actor, '/api/dashboard/reconciliation/', expected=403)
            self.get(actor, '/api/dashboard/reconciliation/audits/', expected=403)
        audit.reason = 'Rewrite'
        with self.assertRaises(ValidationError): audit.save()
        self.assert_preserved(before)

    def test_reconciliation_audit_failure_rolls_back_repair_and_keeps_unassigned_export(self):
        issue = self.issue()
        before = self.preserved()
        with patch('dashboard.reconciliation.WorkflowReconciliationAudit.objects.create', side_effect=DatabaseError('A05 audit outage')):
            with self.assertRaises(DatabaseError): self.repair(issue)
        self.legacy.refresh_from_db()
        self.assertIsNone(self.legacy.academic_semester_id)
        self.assertEqual(WorkflowReconciliationAudit.objects.count(), 0)
        rows = self.report(self.office, semester='unassigned')['supervisor']['records']
        self.assertIn(str(self.legacy.pk), [row['recordId'] for row in rows])
        self.assert_preserved(before)

    def test_changed_reconciliation_evidence_invalidates_a_real_preview(self):
        issue = self.issue()
        self.legacy.status = 'CANCELLED_BY_OFFICE'
        self.legacy.cancellation_reason = 'A05 cancellation after preview'
        self.legacy.save(update_fields=['status', 'cancellation_reason'])
        fresh = self.issue()
        self.assertNotEqual(issue['fingerprint'], fresh['fingerprint'])
        self.assertEqual(self.repair(issue).status_code, 409)
        self.legacy.refresh_from_db()
        self.assertIsNone(self.legacy.academic_semester_id)
        self.assertEqual(WorkflowReconciliationAudit.objects.count(), 0)
        for actor in (self.student_user, self.supervisor, self.coordinator):
            self.client.force_authenticate(actor)
            response = self.client.post(f"/api/dashboard/reconciliation/issues/{fresh['issueId']}/apply/", {
                'expectedFingerprint': fresh['fingerprint'], 'reason': 'Unauthorized repair',
                'resolution': {'action': 'ASSIGN_SEMESTER', 'semesterId': self.semester.pk},
            }, format='json')
            self.assertEqual(response.status_code, 403)
        self.legacy.refresh_from_db()
        self.assertIsNone(self.legacy.academic_semester_id)
        self.assertEqual(WorkflowReconciliationAudit.objects.count(), 0)
