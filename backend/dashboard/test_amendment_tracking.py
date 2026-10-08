from datetime import timedelta
from importlib.util import find_spec
from io import BytesIO
from types import SimpleNamespace

from django.test import SimpleTestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from appointments import test_appointment_lifecycle as fixture


class AmendmentWaitingTests(SimpleTestCase):
    def test_responsible_stage_age_resets_and_terminal_outcomes_stop_waiting(self):
        self.assertIsNotNone(find_spec('dashboard.amendment_tracking'))
        from .amendment_tracking import amendment_waiting_metadata
        now = timezone.now()
        row = SimpleNamespace(status='PENDING_DESTINATION_COORDINATOR',
                              updated_at=now - timedelta(days=3))
        metadata = amendment_waiting_metadata(row, now=now)
        self.assertEqual(metadata['waitingDays'], 3)
        self.assertEqual(metadata['waitingOn'], 'DESTINATION_COORDINATOR')
        row.status = 'APPROVED'
        self.assertIsNone(amendment_waiting_metadata(row, now=now)['waitingDays'])


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class AmendmentTrackingTests(APITestCase):
    def setUp(self):
        fixture.AppointmentLifecycleTests.setUp(self)
        self.office.is_staff = True
        self.office.save(update_fields=['is_staff'])
    _lecturer = fixture.AppointmentLifecycleTests._lecturer

    def request_amendment(self):
        self.client.force_authenticate(self.student_user)
        response = self.client.post('/api/appointments/research-amendments/', {
            'kind': 'RESEARCH', 'title': 'Amended title', 'reason': 'Private research reason',
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        return response.data['id']

    def test_actions_report_dossier_and_export_track_only_authorised_amendments(self):
        pk = self.request_amendment()
        from .actions import build_dashboard_tasks
        self.assertIn(f'research_amendment_{pk}', [r['id'] for r in build_dashboard_tasks(self.supervisor)])
        self.assertNotIn(f'research_amendment_{pk}', [r['id'] for r in build_dashboard_tasks(self.new_supervisor)])
        self.client.force_authenticate(self.office)
        summary = self.client.get('/api/dashboard/summary/')
        self.assertEqual(summary.data['pendingResearchAmendments'], 1)
        report = self.client.get('/api/dashboard/reports/?semester=all')
        self.assertEqual(report.data['researchAmendments']['records'][0]['recordId'], str(pk))
        self.assertEqual(report.data['researchAmendments']['waitingOwnerCounts'], {'SUPERVISOR': 1})
        dossier = self.client.get(f'/api/dashboard/progress/{self.student.matric_no}/')
        self.assertEqual(dossier.data['researchAmendments'][0]['id'], pk)
        exported = self.client.get('/api/dashboard/reports/export/?semester=all')
        workbook = load_workbook(BytesIO(exported.content), read_only=True)
        self.assertIn('Research Amendments', workbook.sheetnames)
        self.assertNotIn('Private research reason', str(list(workbook['Research Amendments'].values)))
        self.client.force_authenticate(self.new_supervisor)
        self.assertEqual(self.client.get('/api/dashboard/reports/?semester=all').data['researchAmendments']['records'], [])

    def test_programme_mismatch_and_stale_baseline_require_review_without_repair(self):
        pk = self.request_amendment()
        from appointments.models import StudentResearchProfile
        from .amendment_tracking import amendment_reconciliation_issues
        StudentResearchProfile.objects.filter(pk=self.profile.pk).update(programme='MASTER OF SOFTWARE ENGINEERING', revision=9)
        issues = amendment_reconciliation_issues()
        kinds = {row.issue_type for row in issues}
        self.assertIn('RESEARCH_PROGRAMME_MISMATCH', kinds)
        self.assertIn('RESEARCH_AMENDMENT_STALE', kinds)
        self.assertIn('RESEARCH_REVISION_CHAIN_BROKEN', kinds)
        self.assertTrue(all(row.repairability == 'REVIEW_REQUIRED' and not row.suggestion for row in issues))

    def test_supporting_supervisor_sees_current_research_without_private_amendments(self):
        self.request_amendment()
        from appointments.models import CoSupervisorAppointment, CoSupervisorNomination
        nomination = CoSupervisorNomination.objects.create(
            student=self.student, primary_appointment=self.supervisor_appointment,
            nominator=self.supervisor, candidate=self.new_supervisor,
            academic_semester=self.semester, justification='Research expertise', status='APPROVED',
        )
        CoSupervisorAppointment.objects.create(
            nomination=nomination, student=self.student,
            supervisor=self.new_supervisor, approved_by=self.coordinator,
        )
        self.client.force_authenticate(self.new_supervisor)
        response = self.client.get(f'/api/dashboard/progress/{self.student.matric_no}/')
        self.assertEqual(response.status_code, 404, response.data)
        response = self.client.get(f'/api/appointments/co-supervisor/students/{self.student.pk}/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertNotIn('Private research reason', str(response.data))

    def test_source_coordinator_loses_dossier_access_after_transfer(self):
        self.client.force_authenticate(self.office)
        response = self.client.post('/api/appointments/research-amendments/', {
            'kind': 'TRANSFER', 'studentId': self.student.pk,
            'destinationProgramme': 'MASTER OF SOFTWARE ENGINEERING', 'reason': 'New programme',
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        pk = response.data['id']
        for actor in (self.coordinator, self.other_coordinator):
            self.client.force_authenticate(actor)
            stage = self.client.get(f'/api/appointments/research-amendments/{pk}/').data['status']
            result = self.client.post(f'/api/appointments/research-amendments/{pk}/decision/',
                                      {'decision': 'APPROVE', 'retainTeam': True, 'expectedStatus': stage}, format='json')
            self.assertEqual(result.status_code, 200, result.data)
        self.client.force_authenticate(self.coordinator)
        self.assertEqual(self.client.get(f'/api/dashboard/progress/{self.student.matric_no}/').status_code, 404)

    def test_latest_revision_must_match_current_research_values(self):
        self.client.force_authenticate(self.office)
        response = self.client.post('/api/appointments/research-amendments/corrections/', {
            'studentId': self.student.pk, 'title': 'Corrected title', 'reason': 'Spelling',
            'meaningUnchanged': True, 'expectedRevision': 0,
        }, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        from appointments.models import StudentResearchProfile
        from .amendment_tracking import amendment_reconciliation_issues
        StudentResearchProfile.objects.filter(pk=self.profile.pk).update(proposed_topic='Unaudited external edit')
        self.assertIn('RESEARCH_REVISION_CHAIN_BROKEN', {row.issue_type for row in amendment_reconciliation_issues()})
