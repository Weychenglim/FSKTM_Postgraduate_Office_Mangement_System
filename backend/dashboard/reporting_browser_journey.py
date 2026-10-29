"""Opt-in A05 production-component/real-HTTP browser acceptance."""
from io import BytesIO
import json
from pathlib import Path
import subprocess
from urllib.request import Request, urlopen

from django.test import LiveServerTestCase, override_settings
from openpyxl import load_workbook
from rest_framework_simplejwt.tokens import AccessToken

from appointments.models import SupervisorApplication
from .acceptance_fixtures import ReportingAcceptanceFixture
from .models import SemesterTimelineEntry, TimelineAuditLog, WorkflowReconciliationAudit


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class ReportingBrowserJourney(ReportingAcceptanceFixture, LiveServerTestCase):
    host = '127.0.0.1'

    def run_journey(self, width, mode):
        from django.apps import apps
        models = [model for model in apps.get_models() if model._meta.app_label in
                  {'accounts', 'academics', 'appointments', 'marks', 'dashboard'}]
        before = {model: list(model.objects.order_by('pk').values()) for model in models}
        actors = {'office': self.office, 'coordinator': self.coordinator, 'primary': self.supervisor,
                  'supporting': self.new_supervisor, 'student': self.student_user, 'foreign': self.foreign_user}
        config = {
            'backend': self.live_server_url, 'width': width, 'mode': mode,
            'tokens': {key: str(AccessToken.for_user(actor)) for key, actor in actors.items()},
            'matricNo': self.student.matric_no, 'studentName': self.student_user.full_name,
            'foreignMatric': self.foreign_student.matric_no, 'foreignName': self.foreign_user.full_name,
            'programme': self.student.programme, 'foreignProgramme': self.foreign_student.programme,
            'semesterId': self.semester.pk, 'semesterCode': self.semester.code,
            'legacyId': self.legacy.pk, 'publicEntryId': self.public_event.pk,
            'startDate': self.public_event.deadline_start.isoformat(),
            'endDate': self.public_event.deadline_end.isoformat(),
        }
        completed = subprocess.run(['node', 'tests/reportingJourney.browser.mjs'],
            cwd=Path(__file__).resolve().parents[2] / 'frontend', input=json.dumps(config),
            capture_output=True, text=True, timeout=300)
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        print(completed.stdout.strip())
        if mode == 'reconciliation':
            self.legacy.refresh_from_db()
            self.assertEqual(self.legacy.academic_semester_id, self.semester.pk)
            audit = WorkflowReconciliationAudit.objects.get()
            self.assertEqual(audit.actor_id, self.office.pk)
            self.assertEqual(audit.reason, 'A05 verified browser repair')
            self.assertIsNone(audit.before_values['academicSemesterId'])
            self.assertEqual(audit.after_values['academicSemesterId'], self.semester.pk)
            self.assertEqual(list(TimelineAuditLog.objects.order_by('id').values_list('action', flat=True)),
                             ['ADD_ENTRY', 'EDIT_ENTRY', 'DELETE_ENTRY'])
            self.assertEqual(set(TimelineAuditLog.objects.values_list('actor_id', flat=True)), {self.office.pk})
            self.assertFalse(SemesterTimelineEntry.objects.filter(title='A05 edited browser event').exists())
            expected_applications = [dict(row) for row in before[SupervisorApplication]]
            for row in expected_applications:
                if row['id'] == self.legacy.pk:
                    row['academic_semester_id'] = self.semester.pk
            self.assertEqual(list(SupervisorApplication.objects.order_by('pk').values()), expected_applications)
            allowed = {SupervisorApplication, WorkflowReconciliationAudit, TimelineAuditLog}
        else:
            allowed = set()
        for model, rows in before.items():
            if model not in allowed:
                self.assertEqual(list(model.objects.order_by('pk').values()), rows, model.__name__)
        # Parse actual downloaded bytes independently of the browser's UI assertions.
        for key in ('office', 'coordinator', 'primary', 'supporting'):
            headers = {'Authorization': 'Bearer ' + config['tokens'][key]}
            with urlopen(Request(self.live_server_url + '/api/dashboard/reports/export/?semester=all', headers=headers)) as response:
                workbook = load_workbook(BytesIO(response.read()))
            if key == 'coordinator':
                self.assertNotIn('Marks', workbook.sheetnames)
                self.assertNotIn(self.foreign_student.matric_no, str(list(workbook['Supervisor'].values)))
            self.assertFalse(any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row))

    def test_reports_desktop(self): self.run_journey(1440, 'reports')
    def test_reports_phone(self): self.run_journey(390, 'reports')
    def test_reconciliation_desktop(self): self.run_journey(1440, 'reconciliation')
    def test_reconciliation_phone(self): self.run_journey(390, 'reconciliation')
    def test_dossier_race_desktop(self): self.run_journey(1440, 'race')
    def test_dossier_race_phone(self): self.run_journey(390, 'race')


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class TerminalLifecycleBrowserJourney(ReportingAcceptanceFixture, LiveServerTestCase):
    """A06 terminal privacy through real HTTP, separate from historical A05 counts."""
    host = '127.0.0.1'
    run_journey = ReportingBrowserJourney.run_journey

    def setUp(self):
        from accounts.participant_lifecycle import transition_student
        super().setUp()
        transition_student(matric_no=self.student.matric_no, actor=self.office,
            target_status='WITHDRAWN', reason='PRIVATE lifecycle decision')

    def test_terminal_privacy_desktop(self): self.run_journey(1440, 'privacy')
    def test_terminal_privacy_phone(self): self.run_journey(390, 'privacy')
